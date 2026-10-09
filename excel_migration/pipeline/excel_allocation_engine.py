"""
excel_allocation_engine.py
Phase 3: Recreates BOM matching and stock calculation logic for the Excel app.
Imports the untouched Streamlit allocation_engine.py to preserve exact business rules.
Uses Python for the complex FIFO loop and incomplete BOM detection.
Uses Excel formulas for True Current Stock and Warning calculations.
"""
import os
import sys
import pandas as pd
import sqlite3
import openpyxl
from openpyxl.utils.dataframe import dataframe_to_rows

import config
import excel_data_loader as dl
import refresh_data as rd

# Add Streamlit app directory to path so we can import allocation_engine directly
APP_DIR = r"d:\Planner Dashboard"
if APP_DIR not in sys.path:
    sys.path.append(APP_DIR)
import allocation_engine as ae

def get_tcf_drops(files_dict):
    """Combines VGL1 and VGL2 into a single tcf_drops dataframe."""
    dfs = []
    if files_dict.get('VGL1'):
        df1 = dl.load_vgl(files_dict['VGL1'])
        if df1 is not None and not df1.empty: dfs.append(df1)
    if files_dict.get('VGL2'):
        df2 = dl.load_vgl(files_dict['VGL2'])
        if df2 is not None and not df2.empty: dfs.append(df2)
    if not dfs:
        return pd.DataFrame()
    return pd.concat(dfs, ignore_index=True)

def load_sqlite_data():
    """Loads master data from clear_to_build.db"""
    db_path = os.path.join(APP_DIR, "clear_to_build.db")
    conn = sqlite3.connect(db_path)
    
    bom = pd.read_sql("SELECT short_vehicle_code as 'Short Vehicle Code', front_wiring as 'Front Wiring', cockpit as 'Cockpit', engine as 'Engine' FROM bom_details", conn)
    
    engine_stocks = pd.read_sql("SELECT tcf_line as 'TCF Line', engine_part_no as 'Engine Part No', model as 'Model', ta_code as 'TA Code', clearance_qty as 'Clearance After 6:30AM' FROM engine_stocks", conn)
    
    nova_stocks = pd.read_sql("SELECT * FROM nova_stocks", conn)
    
    try:
        model_shortages = pd.read_sql("SELECT * FROM model_shortages", conn)
    except:
        model_shortages = pd.DataFrame()
        
    conn.close()
    return bom, engine_stocks, nova_stocks, model_shortages

def write_table_data(ws, df, start_row, formulas=None):
    """Writes dataframe to an Excel table area, injecting formulas where provided."""
    # Delete existing rows
    ws.delete_rows(start_row, max(ws.max_row, start_row + 1))
    
    if df is None or df.empty: return
    
    # Replace NaN with None
    df = df.where(pd.notnull(df), None)
    
    for r_idx, row in enumerate(dataframe_to_rows(df, index=False, header=False), start=start_row):
        for c_idx, value in enumerate(row, start=1):
            if formulas and c_idx in formulas:
                ws.cell(row=r_idx, column=c_idx, value=formulas[c_idx])
            else:
                ws.cell(row=r_idx, column=c_idx, value=value)

def run_phase3():
    print("Starting Phase 3 Allocation Engine...")
    settings = config.load_settings()
    data_folder = settings.get("Local Data Folder", r"D:\Dashboard Files")
    
    # 1. Load Data
    files = {
        'FLOAT': rd.get_latest_file(data_folder, settings.get('Float Report Pattern', 'PPC_Float_Report_*'), exclude='Paint'),
        'VGL1': rd.get_latest_file(data_folder, settings.get('TCF1 VGL Pattern', 'DPT_PLAN-VIN_GENERATION_REPORT_*')),
        'VGL2': rd.get_latest_file(data_folder, settings.get('TCF2 VGL Pattern', 'TCF2_DPT-PLAN_VIN_GENERATION_REPORT_*')),
    }
    
    # Load SQLite master data
    bom, engine_stocks, nova_stocks, model_shortages = load_sqlite_data()
    
    # Load grouped stocks (Cockpit & Wiring)
    part_ck_tcf1_nova, _ = dl.load_stock_grouped(
        rd.get_latest_file(data_folder, settings.get('Cockpit Nova Pattern', 'Fresh VIN PPC*')),
        settings.get('Cockpit Nova Sheet', 'Fresh VIN PPC'),
        int(settings.get('Cockpit Nova VC Col', 1)),
        int(settings.get('Cockpit Nova Part Col', 0)),
        int(settings.get('Cockpit Nova Qty Col', 9)),
        int(settings.get('Cockpit Nova Skip', 3))
    )
    
    part_ck_tcf1_altroz, _ = dl.load_stock_grouped(
        rd.get_latest_file(data_folder, settings.get('Cockpit Altroz Pattern', 'Fresh VIN PPC*')),
        settings.get('Cockpit Altroz Sheet', 'Fresh VIN PPC'),
        int(settings.get('Cockpit Altroz VC Col', 4)),
        int(settings.get('Cockpit Altroz Part Col', 3)),
        int(settings.get('Cockpit Altroz Qty Col', 12)),
        int(settings.get('Cockpit Altroz Skip', 3))
    )
    
    part_ck_tcf2, _ = dl.load_stock_grouped(
        rd.get_latest_file(data_folder, settings.get('Cockpit TCF2 Pattern', 'Fresh vin PPC*')),
        settings.get('Cockpit TCF2 Sheet', 'Fresh vin PPC'),
        int(settings.get('Cockpit TCF2 VC Col', 3)),
        int(settings.get('Cockpit TCF2 Part Col', 1)),
        int(settings.get('Cockpit TCF2 Qty Col', 9)),
        int(settings.get('Cockpit TCF2 Skip', 3))
    )
    
    part_wh_tcf1, _ = dl.load_stock_grouped(
        rd.get_latest_file(data_folder, settings.get('Wiring TCF1 Pattern', 'Coverage file 6.30 AM New*')),
        settings.get('Wiring TCF1 Sheet', 'Coverage file 6.30 AM New'),
        int(settings.get('Wiring TCF1 VC Col', 2)),
        int(settings.get('Wiring TCF1 Part Col', 3)),
        int(settings.get('Wiring TCF1 Qty Col', 9)),
        int(settings.get('Wiring TCF1 Skip', 2))
    )
    
    part_wh_tcf2, _ = dl.load_stock_grouped(
        rd.get_latest_file(data_folder, settings.get('Wiring TCF2 Pattern', 'coverage file 6.30 pm*')),
        settings.get('Wiring TCF2 Sheet', 'coverage file 6.30 pm'),
        int(settings.get('Wiring TCF2 VC Col', 1)),
        int(settings.get('Wiring TCF2 Part Col', 2)),
        int(settings.get('Wiring TCF2 Qty Col', 9)),
        int(settings.get('Wiring TCF2 Skip', 2))
    )
    
    # Merge stocks
    cockpit_shift_start = {**part_ck_tcf1_nova, **part_ck_tcf1_altroz, **part_ck_tcf2}
    wiring_shift_start = {**part_wh_tcf1, **part_wh_tcf2}
    
    engine_shift_start = {}
    for _, r in engine_stocks.iterrows():
        engine_shift_start[str(r['Engine Part No']).strip()] = int(r['Clearance After 6:30AM'])
        
    nova_shift_start = {}
    for _, r in nova_stocks.iterrows():
        nova_shift_start[str(r['Material']).strip()] = int(r['Clearance Qty'])
        
    tcf_drops = get_tcf_drops(files)
    df_float = dl.load_float_report(files['FLOAT']) if files['FLOAT'] else pd.DataFrame()
    
    # 2. Run Python Allocation Engine for True Stock
    # Using exact Streamlit business logic
    _, engine_cons, _ = ae.calculate_true_stock(engine_shift_start, tcf_drops, bom, 'Engine')
    _, nova_cons, _ = ae.calculate_true_stock(nova_shift_start, tcf_drops, bom, 'Engine')
    _, cockpit_cons, _ = ae.calculate_true_stock(cockpit_shift_start, tcf_drops, bom, 'Cockpit')
    _, wiring_cons, _ = ae.calculate_true_stock(wiring_shift_start, tcf_drops, bom, 'Front Wiring')
    
    # 3. Write Excel Tables with Embedded Formulas for True Stock & Warnings
    wb = openpyxl.load_workbook(config.EXCEL_PATH, keep_vba=True)
    
    # BOM
    print("Writing BOM...")
    write_table_data(wb['BOM'], bom, 3)
    
    # Engine Stock
    print("Writing Engine Stocks...")
    eng_rows = []
    for p, q in engine_shift_start.items():
        eng_rows.append({
            'Part Number': p, 'Part Type': 'Engine', 'Shift Start Stock': q, 
            'Consumed': engine_cons.get(p, 0),
            'True Current Stock': None, 'Virtual Stock': 0, 'Status': '', 'Warning': None
        })
    df_eng = pd.DataFrame(eng_rows)
    eng_formulas = {
        5: "=[@[Shift Start Stock]]-[@Consumed]",
        8: '=IF([@[True Current Stock]]<0, "⚠️ Negative stock (backflush)", "")'
    }
    write_table_data(wb['Stock'], df_eng, 4, eng_formulas)
    
    # Nova Stock
    print("Writing Nova Stocks...")
    nova_rows = []
    for p, q in nova_shift_start.items():
        nova_rows.append({
            'Material': p, 'Shift Start Stock': q, 'Consumed': nova_cons.get(p, 0),
            'True Current Stock': None, 'Virtual Stock': 0, 'Warning': None
        })
    df_nova = pd.DataFrame(nova_rows)
    nova_formulas = {
        4: "=[@[Shift Start Stock]]-[@Consumed]",
        6: '=IF([@[True Current Stock]]<0, "⚠️ Negative stock (backflush)", "")'
    }
    write_table_data(wb['Stock'], df_nova, 20, nova_formulas) # Offset for tblNovaStock
    
    # Cockpit Stock
    print("Writing Cockpit Stocks...")
    ck_rows = []
    for p, q in cockpit_shift_start.items():
        ck_rows.append({
            'Part Number': p, 'Shift Start Stock': q, 'Consumed': cockpit_cons.get(p, 0),
            'True Current Stock': None, 'Virtual Stock': 0, 'Status': '', 'Warning': None
        })
    df_ck = pd.DataFrame(ck_rows)
    ck_formulas = {
        4: "=[@[Shift Start Stock]]-[@Consumed]",
        7: '=IF([@[True Current Stock]]<0, "⚠️ Negative stock (backflush)", "")'
    }
    write_table_data(wb['Cockpit'], df_ck, 3, ck_formulas)
    
    # Wiring Stock
    print("Writing Wiring Stocks...")
    wh_rows = []
    for p, q in wiring_shift_start.items():
        wh_rows.append({
            'Part Number': p, 'Shift Start Stock': q, 'Consumed': wiring_cons.get(p, 0),
            'True Current Stock': None, 'Virtual Stock': 0, 'Status': '', 'Warning': None
        })
    df_wh = pd.DataFrame(wh_rows)
    write_table_data(wb['Front Wiring'], df_wh, 3, ck_formulas) # Same columns as cockpit
    
    # 4. Run FIFO Allocation
    print("Running FIFO Allocation Engine...")
    true_engine = {p: q - engine_cons.get(p, 0) for p, q in engine_shift_start.items()}
    true_nova = {p: q - nova_cons.get(p, 0) for p, q in nova_shift_start.items()}
    true_cockpit = {p: q - cockpit_cons.get(p, 0) for p, q in cockpit_shift_start.items()}
    true_wiring = {p: q - wiring_cons.get(p, 0) for p, q in wiring_shift_start.items()}
    
    ms_list = model_shortages.to_dict('records') if not model_shortages.empty else []
    
    allocation_results, _ = ae.run_allocation(df_float, bom, true_engine, true_cockpit, true_wiring, true_nova, ms_list)
    df_alloc = pd.DataFrame(allocation_results)
    
    print("Writing FIFO Allocation...")
    write_table_data(wb['FIFO Allocation'], df_alloc, 3)
    
    # Log
    wb['Log'].append([pd.Timestamp.now().strftime("%d-%m-%Y %I:%M %p"), "Allocation Engine", "System", f"Ran FIFO allocation for {len(df_alloc)} cabs.", "✅"])
    
    wb.save(config.EXCEL_PATH)
    print("Phase 3 complete.")

if __name__ == "__main__":
    run_phase3()
