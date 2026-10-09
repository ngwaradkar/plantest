import streamlit as st
import pandas as pd
import altair as alt
import numpy as np
import io
from datetime import datetime, timedelta
import pytz
import os
import sys

# Add project root to sys.path
sys.path.append(os.path.abspath(os.path.join(os.path.dirname(__file__), '..')))
import data_loader as dl
import allocation_engine as ae

def render():
    db_folder_path = st.session_state.get('_config', {}).get('db_folder_path', r'C:\PPC_DB')
    try:
        shop_totals, shop_vehicles_df, hourly_df = dl.load_shop_wise_report(db_folder_path)
    except Exception:
        shop_totals, shop_vehicles_df, hourly_df = None, None, None
        
    classified_files = st.session_state.get("classified_files", {})
    from core.data_manager import get_dashboard_data, ENGINE_DEFAULT_DATA
    engine_default_data = ENGINE_DEFAULT_DATA
    try:
        data = get_dashboard_data(classified_files)
        paint_summary_dict = data.paint_summary if (data and data.status != "error") else {}
        paint_summary_vc_dict = data.paint_summary_vc_dict if (data and data.status != "error") else {}
        float_df = data.float_df if (data and data.status != "error") else pd.DataFrame()
        tcf1_drops = data.tcf1_drops
        tcf2_drops = data.tcf2_drops
        tcf1_alloc_df = data.tcf1_alloc_df
        tcf2_alloc_df = data.tcf2_alloc_df
        bom_df = data.bom_df
        tcf1_cockpit_start = data.true_cockpit_tcf1
        tcf2_cockpit_start = data.true_cockpit_tcf2
        tcf1_wiring_start = data.true_wiring_tcf1
        tcf2_wiring_start = data.true_wiring_tcf2
    except Exception:
        paint_summary_dict = {}
        paint_summary_vc_dict = {}
        float_df = pd.DataFrame()
        tcf1_drops = None
        tcf2_drops = None
        tcf1_alloc_df = pd.DataFrame()
        tcf2_alloc_df = pd.DataFrame()
        bom_df = pd.DataFrame()
        tcf1_cockpit_start = {}
        tcf2_cockpit_start = {}
        tcf1_wiring_start = {}
        tcf2_wiring_start = {}
        
    temp_float_df = float_df.copy() if not float_df.empty else pd.DataFrame()
        
    # --- SECTION 0: SHOP-WISE PLANT PRODUCTION SUMMARY ---
    if shop_totals is not None or shop_vehicles_df is not None:
        st.markdown("### 🏭 Shop-Wise Plant Production Summary (Daily Report)")
        if shop_totals:
            rep_date = shop_totals.get('Date', shop_totals.get('REPORT DATE', '03/08/2026'))
            cap_col1, cap_col2 = st.columns([2, 2])
            with cap_col1:
                st.caption(f"📅 **Report Date**: {rep_date}")
            with cap_col2:
                last_gen = st.session_state.get('last_generated_at')
                if last_gen:
                    st.caption(f"🕒 **Last Generated**: {last_gen.strftime('%d-%b-%Y %I:%M %p')}")

            # Sticky KPI bar: stays pinned to the top while scrolling through the
            # rest of the (often long) summary report below.
            st.markdown("""
                <style>
                .st-key-sticky_kpi_bar {
                    position: sticky;
                    top: 2.75rem;
                    z-index: 998;
                    background-color: var(--background-color, #0e1117);
                    padding: 0.5rem 0 0.75rem 0;
                    border-bottom: 1px solid rgba(128,128,128,0.25);
                }
                </style>
            """, unsafe_allow_html=True)

            with st.container(key="sticky_kpi_bar"):
                s_kpi1, s_kpi2, s_kpi3, s_kpi4, s_kpi5, s_kpi6 = st.columns(6)
                s_kpi1.metric("TCF1 VIN Count", f"{shop_totals.get('TCF VIN', 0)} cabs")
                s_kpi2.metric("TCF2 VIN Count", f"{shop_totals.get('TCF2 VIN', 0)} cabs")
                s_kpi3.metric("Total TCF Dropping", f"{int(shop_totals.get('TCF DROP', 0)) + int(shop_totals.get('TCF2 DROP', 0))} cabs")
                s_kpi4.metric("Paint Lifting", f"{shop_totals.get('PAINT', 0)} cabs")
                s_kpi5.metric("T60 Count", f"{shop_totals.get('T60', 0)} cabs")
                s_kpi6.metric("T40 Count", f"{shop_totals.get('T40', 0)} cabs")
            
        if shop_vehicles_df is not None and not shop_vehicles_df.empty:
            st.markdown("#### 🚗 Model-Wise Production Matrix (TCF1 & TCF2 Breakdown)")
            
            tcf1_models = ['PUNCH', 'PUNCH Exports', 'PUNCH EV', 'ALTROZ', 'ALTROZ DCA', 'ALTROZ EV']
            tcf2_models = ['HARRIER EV', 'SAFARI', 'HARRIER', 'SAFARI EV']
            
            df1 = shop_vehicles_df[shop_vehicles_df['Model'].isin(tcf1_models)].copy()
            df2 = shop_vehicles_df[shop_vehicles_df['Model'].isin(tcf2_models)].copy()
            
            t1_vin = int(df1['VIN'].sum()) if not df1.empty else 0
            t1_drop = int(df1['Drop'].sum()) if not df1.empty else 0
            t1_paint = int(df1['Paint Lifting'].sum()) if not df1.empty else 0
            t1_t60 = int(df1['T60'].sum()) if not df1.empty else 0
            t1_t40 = int(df1['T40'].sum()) if not df1.empty else 0

            t2_vin = int(df2['VIN'].sum()) if not df2.empty else 0
            t2_drop = int(df2['Drop'].sum()) if not df2.empty else 0
            t2_paint = int(df2['Paint Lifting'].sum()) if not df2.empty else 0
            t2_t60 = int(df2['T60'].sum()) if not df2.empty else 0
            t2_t40 = int(df2['T40'].sum()) if not df2.empty else 0

            g_vin = t1_vin + t2_vin
            g_drop = t1_drop + t2_drop
            g_paint = t1_paint + t2_paint
            g_t60 = t1_t60 + t2_t60
            g_t40 = t1_t40 + t2_t40

            html_table = f"""<style>
.matrix-card {{
    background: rgba(15, 23, 42, 0.02);
    border: 1px solid rgba(226, 232, 240, 0.9);
    border-radius: 14px;
    padding: 12px;
    margin-top: 8px;
    margin-bottom: 16px;
    box-shadow: 0 10px 25px -5px rgba(0, 0, 0, 0.05);
}}
.matrix-table {{
    width: 100%;
    border-collapse: separate;
    border-spacing: 0;
    border-radius: 10px;
    overflow: hidden;
    font-family: system-ui, -apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, sans-serif;
}}
.matrix-table th {{
    background: linear-gradient(135deg, #0f172a 0%, #1e293b 100%);
    color: #38bdf8;
    padding: 12px 16px;
    font-size: 13px;
    font-weight: 700;
    text-transform: uppercase;
    letter-spacing: 0.6px;
    text-align: center;
    border-bottom: 2px solid #0284c7;
}}
.matrix-table th:first-child {{ text-align: left; padding-left: 20px; }}
.matrix-table td {{
    padding: 11px 16px;
    font-size: 14px;
    color: #1e293b;
    text-align: center;
    border-bottom: 1px solid #e2e8f0;
    font-weight: 600;
}}
.matrix-table td:first-child {{ text-align: left; padding-left: 20px; }}
.tr-tcf1 {{ background-color: #ffffff; }}
.tr-tcf1:hover {{ background-color: #f0f9ff; }}
.tr-tcf2 {{ background-color: #ffffff; }}
.tr-tcf2:hover {{ background-color: #f0fdf4; }}
.tr-tcf1-tot td {{
    background: linear-gradient(90deg, #1e40af 0%, #3b82f6 100%) !important;
    color: #ffffff !important;
    font-weight: 800 !important;
    font-size: 14px !important;
    border-top: 2px solid #1d4ed8 !important;
    border-bottom: 2px solid #1d4ed8 !important;
}}
.tr-tcf2-tot td {{
    background: linear-gradient(90deg, #0f766e 0%, #14b8a6 100%) !important;
    color: #ffffff !important;
    font-weight: 800 !important;
    font-size: 14px !important;
    border-top: 2px solid #0d9488 !important;
    border-bottom: 2px solid #0d9488 !important;
}}
.tr-grand-tot td {{
    background: linear-gradient(90deg, #312e81 0%, #4f46e5 100%) !important;
    color: #fbbf24 !important;
    font-weight: 900 !important;
    font-size: 15px !important;
    letter-spacing: 0.5px;
    border-top: 3px solid #6366f1 !important;
}}
.badge-tcf1 {{
    background: #dbeafe;
    color: #1e40af;
    padding: 3px 8px;
    border-radius: 6px;
    font-size: 11px;
    font-weight: 700;
    display: inline-block;
}}
.badge-tcf2 {{
    background: #ccfbf1;
    color: #0f766e;
    padding: 3px 8px;
    border-radius: 6px;
    font-size: 11px;
    font-weight: 700;
    display: inline-block;
}}
</style>
<div class="matrix-card">
<table class="matrix-table">
<thead>
<tr>
<th>Model</th>
<th>VIN</th>
<th>Drop</th>
<th>Paint Lifting</th>
<th>T60</th>
<th>T40</th>
</tr>
</thead>
<tbody>"""
            
            # TCF1 Rows
            if not df1.empty:
                for _, r in df1.iterrows():
                    html_table += f"""<tr class="tr-tcf1"><td><span class="badge-tcf1">TCF1</span> &nbsp; <b>{r['Model']}</b></td><td>{r['VIN']}</td><td>{r['Drop']}</td><td>{r['Paint Lifting']}</td><td>{r['T60']}</td><td>{r['T40']}</td></tr>"""
            
            # TCF1 Total
            html_table += f"""<tr class="tr-tcf1-tot"><td>🔹 TOTAL TCF1</td><td>{t1_vin}</td><td>{t1_drop}</td><td>{t1_paint}</td><td>{t1_t60}</td><td>{t1_t40}</td></tr>"""

            # TCF2 Rows
            if not df2.empty:
                for _, r in df2.iterrows():
                    html_table += f"""<tr class="tr-tcf2"><td><span class="badge-tcf2">TCF2</span> &nbsp; <b>{r['Model']}</b></td><td>{r['VIN']}</td><td>{r['Drop']}</td><td>{r['Paint Lifting']}</td><td>{r['T60']}</td><td>{r['T40']}</td></tr>"""
            
            # TCF2 Total
            html_table += f"""<tr class="tr-tcf2-tot"><td>🔸 TOTAL TCF2</td><td>{t2_vin}</td><td>{t2_drop}</td><td>{t2_paint}</td><td>{t2_t60}</td><td>{t2_t40}</td></tr>"""

            # Grand Total
            html_table += f"""<tr class="tr-grand-tot"><td>🏆 GRAND TOTAL PLANT</td><td>{g_vin}</td><td>{g_drop}</td><td>{g_paint}</td><td>{g_t60}</td><td>{g_t40}</td></tr></tbody></table></div>"""
            
            st.markdown(html_table, unsafe_allow_html=True)
            
            # Excel export button for Model Wise Matrix
            df1_sub = df1.copy()
            t1_row = pd.DataFrame([{'Model': 'TOTAL TCF1', 'VIN': t1_vin, 'Drop': t1_drop, 'Paint Lifting': t1_paint, 'T60': t1_t60, 'T40': t1_t40}])
            df2_sub = df2.copy()
            t2_row = pd.DataFrame([{'Model': 'TOTAL TCF2', 'VIN': t2_vin, 'Drop': t2_drop, 'Paint Lifting': t2_paint, 'T60': t2_t60, 'T40': t2_t40}])
            gt_row = pd.DataFrame([{'Model': 'GRAND TOTAL PLANT', 'VIN': g_vin, 'Drop': g_drop, 'Paint Lifting': g_paint, 'T60': g_t60, 'T40': g_t40}])
            
            export_matrix_df = pd.concat([df1_sub, t1_row, df2_sub, t2_row, gt_row], ignore_index=True)
            
            buf_matrix = io.BytesIO()
            with pd.ExcelWriter(buf_matrix, engine='openpyxl') as writer:
                export_matrix_df.to_excel(writer, index=False, sheet_name='Production Matrix')
            st.download_button(
                label="📥 Export Model-Wise Production Matrix to Excel",
                data=buf_matrix.getvalue(),
                file_name="Model_Wise_Production_Matrix.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                key="export_prod_matrix_btn"
            )
            
        st.markdown("---")

    # --- SECTION 0.5: HOURLY PRODUCTION & GENERATION TRACKER (TCF1 & TCF2) ---
    if hourly_df is not None and not hourly_df.empty:
        st.markdown("### ⏱️ Hourly Production & Line Generation Tracker")
        st.caption("Live hour-by-hour tracking of VIN Generation and TCF Dropping across TCF1 and TCF2 assembly lines (Format: **Hourly Output [Shift Cumulative Output]**).")
        
        point_col = 'ACHIEVEMENT POINT' if 'ACHIEVEMENT POINT' in hourly_df.columns else hourly_df.columns[0]
        time_cols = [c for c in hourly_df.columns if '-' in str(c) and ('AM' in str(c).upper() or 'PM' in str(c).upper())]
        tot_cols = [c for c in hourly_df.columns if 'TOTAL' in str(c).upper()]
        
        def _parse_cell_val(val):
            if pd.isna(val):
                return 0, 0
            s = str(val).strip()
            m = re.match(r'^(\d+)(?:\[(\d+)\])?$', s)
            if m:
                h = int(m.group(1))
                c = int(m.group(2)) if m.group(2) else h
                return h, c
            try:
                n = int(float(s))
                return n, n
            except Exception:
                return 0, 0

        def _fmt_cell(h, c):
            if c != 0 or h != 0:
                return f"{h}[{c}]"
            return "0[0]"

        rows_by_point = {}
        for _, r in hourly_df.iterrows():
            rows_by_point[str(r[point_col]).strip()] = r
            
        tcf1_vin = rows_by_point.get('TCF1_VIN_GENERATION')
        tcf2_vin = rows_by_point.get('TCF2_VIN_GENERATION')
        tcf1_drop = rows_by_point.get('TCF1-DROP')
        tcf2_drop = rows_by_point.get('TCF2-DROP')
        
        # Calculate Total VIN Generation Row
        tot_vin_row = {point_col: 'TOTAL VIN GENERATION'}
        for c in time_cols:
            h1, c1 = _parse_cell_val(tcf1_vin[c]) if tcf1_vin is not None else (0, 0)
            h2, c2 = _parse_cell_val(tcf2_vin[c]) if tcf2_vin is not None else (0, 0)
            tot_vin_row[c] = _fmt_cell(h1 + h2, c1 + c2)
        for c in tot_cols:
            v1 = int(tcf1_vin[c]) if tcf1_vin is not None and not pd.isna(tcf1_vin[c]) else 0
            v2 = int(tcf2_vin[c]) if tcf2_vin is not None and not pd.isna(tcf2_vin[c]) else 0
            tot_vin_row[c] = v1 + v2

        # Calculate Total TCF Dropping Row
        tot_drop_row = {point_col: 'TOTAL TCF DROPPING'}
        for c in time_cols:
            h1, c1 = _parse_cell_val(tcf1_drop[c]) if tcf1_drop is not None else (0, 0)
            h2, c2 = _parse_cell_val(tcf2_drop[c]) if tcf2_drop is not None else (0, 0)
            tot_drop_row[c] = _fmt_cell(h1 + h2, c1 + c2)
        for c in tot_cols:
            v1 = int(tcf1_drop[c]) if tcf1_drop is not None and not pd.isna(tcf1_drop[c]) else 0
            v2 = int(tcf2_drop[c]) if tcf2_drop is not None and not pd.isna(tcf2_drop[c]) else 0
            tot_drop_row[c] = v1 + v2

        # KPI Metrics Row
        kpi_col1, kpi_col2, kpi_col3, kpi_col4, kpi_col5, kpi_col6 = st.columns(6)
        
        t1_vin_day = int(tcf1_vin['DAY TOTAL']) if tcf1_vin is not None and 'DAY TOTAL' in tcf1_vin else 0
        t1_vin_sh_a = int(tcf1_vin['SHIFT A TOTAL']) if tcf1_vin is not None and 'SHIFT A TOTAL' in tcf1_vin else 0
        t1_vin_sh_b = int(tcf1_vin['SHIFT B TOTAL']) if tcf1_vin is not None and 'SHIFT B TOTAL' in tcf1_vin else 0
        
        t2_vin_day = int(tcf2_vin['DAY TOTAL']) if tcf2_vin is not None and 'DAY TOTAL' in tcf2_vin else 0
        t2_vin_sh_a = int(tcf2_vin['SHIFT A TOTAL']) if tcf2_vin is not None and 'SHIFT A TOTAL' in tcf2_vin else 0
        t2_vin_sh_b = int(tcf2_vin['SHIFT B TOTAL']) if tcf2_vin is not None and 'SHIFT B TOTAL' in tcf2_vin else 0
        
        tot_vin_day = t1_vin_day + t2_vin_day
        
        t1_drop_day = int(tcf1_drop['DAY TOTAL']) if tcf1_drop is not None and 'DAY TOTAL' in tcf1_drop else 0
        t1_drop_sh_a = int(tcf1_drop['SHIFT A TOTAL']) if tcf1_drop is not None and 'SHIFT A TOTAL' in tcf1_drop else 0
        t1_drop_sh_b = int(tcf1_drop['SHIFT B TOTAL']) if tcf1_drop is not None and 'SHIFT B TOTAL' in tcf1_drop else 0
        
        t2_drop_day = int(tcf2_drop['DAY TOTAL']) if tcf2_drop is not None and 'DAY TOTAL' in tcf2_drop else 0
        t2_drop_sh_a = int(tcf2_drop['SHIFT A TOTAL']) if tcf2_drop is not None and 'SHIFT A TOTAL' in tcf2_drop else 0
        t2_drop_sh_b = int(tcf2_drop['SHIFT B TOTAL']) if tcf2_drop is not None and 'SHIFT B TOTAL' in tcf2_drop else 0
        
        tot_drop_day = t1_drop_day + t2_drop_day

        kpi_col1.metric("TCF1 Day VIN", f"{t1_vin_day} cabs", f"A: {t1_vin_sh_a} | B: {t1_vin_sh_b}")
        kpi_col2.metric("TCF2 Day VIN", f"{t2_vin_day} cabs", f"A: {t2_vin_sh_a} | B: {t2_vin_sh_b}")
        kpi_col3.metric("Plant VIN Total", f"{tot_vin_day} cabs", f"TCF1 + TCF2")
        kpi_col4.metric("TCF1 Day Drop", f"{t1_drop_day} cabs", f"A: {t1_drop_sh_a} | B: {t1_drop_sh_b}")
        kpi_col5.metric("TCF2 Day Drop", f"{t2_drop_day} cabs", f"A: {t2_drop_sh_a} | B: {t2_drop_sh_b}")
        kpi_col6.metric("Plant Drop Total", f"{tot_drop_day} cabs", f"TCF1 + TCF2")

        table_rows_display = []
        if tcf1_vin is not None:
            table_rows_display.append(('TCF1_VIN_GENERATION', 'TCF1 VIN GENERATION', 'tcf1-vin', tcf1_vin))
        if tcf2_vin is not None:
            table_rows_display.append(('TCF2_VIN_GENERATION', 'TCF2 VIN GENERATION', 'tcf2-vin', tcf2_vin))
        if tcf1_vin is not None or tcf2_vin is not None:
            table_rows_display.append(('TOTAL_VIN', 'TOTAL VIN GENERATION', 'tot-vin', tot_vin_row))
        if tcf1_drop is not None:
            table_rows_display.append(('TCF1-DROP', 'TCF1 DROP', 'tcf1-drop', tcf1_drop))
        if tcf2_drop is not None:
            table_rows_display.append(('TCF2-DROP', 'TCF2 DROP', 'tcf2-drop', tcf2_drop))
        if tcf1_drop is not None or tcf2_drop is not None:
            table_rows_display.append(('TOTAL_DROP', 'TOTAL TCF DROPPING', 'tot-drop', tot_drop_row))

        all_headers = [point_col] + time_cols + tot_cols

        def _fmt_hdr_html(hdr_str):
            s = str(hdr_str).strip()
            if 'ACHIEVEMENT' in s.upper():
                return "Stage / Line"
            if 'SHIFT' in s.upper():
                m = re.search(r'SHIFT\s*([A-Z])', s, re.I)
                if m:
                    return f"Shift {m.group(1).upper()}<br><span style='font-size:10px;font-weight:700;opacity:0.9;'>Total</span>"
            if 'DAY' in s.upper():
                return "Day Total<br><span style='font-size:10px;font-weight:700;opacity:0.9;'>Plant</span>"
                
            m = re.match(r'(\d{1,2})(?::(\d{2}))?\s*([AP]M)?\s*-\s*(\d{1,2})(?::(\d{2}))?\s*([AP]M)', s, re.I)
            if m:
                h1, m1, p1, h2, m2, p2 = m.groups()
                t1 = f"{h1}:{m1}" if m1 and m1 != '00' else h1
                t2 = f"{h2}:{m2}" if m2 and m2 != '00' else h2
                period = p2 or p1 or ""
                return f"{t1}-{t2}<br><span style='font-size:10px;font-weight:700;opacity:0.9;'>{period.upper()}</span>"
            return s

        html_hourly_table = f"""<style>
.hourly-card {{
    background: rgba(15, 23, 42, 0.02);
    border: 1px solid rgba(226, 232, 240, 0.9);
    border-radius: 12px;
    padding: 8px 10px;
    margin-top: 6px;
    margin-bottom: 14px;
    box-shadow: 0 4px 15px -3px rgba(0, 0, 0, 0.05);
    overflow: hidden;
    width: 100%;
    box-sizing: border-box;
}}
.hourly-table {{
    width: 100%;
    table-layout: fixed;
    border-collapse: separate;
    border-spacing: 0;
    border-radius: 8px;
    overflow: hidden;
    font-family: system-ui, -apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, sans-serif;
    font-size: 12.5px;
}}
.hourly-table th {{
    background: linear-gradient(135deg, #0f172a 0%, #1e293b 100%);
    color: #38bdf8;
    padding: 8px 3px;
    font-size: 12px;
    font-weight: 800;
    text-transform: uppercase;
    letter-spacing: 0.25px;
    text-align: center;
    border-bottom: 2px solid #0284c7;
    line-height: 1.25;
    overflow: hidden;
}}
.hourly-table th:first-child {{ text-align: left; padding-left: 10px; width: 17%; }}
.hourly-table th.th-time {{ width: 7.2%; }}
.hourly-table th.th-shift {{ width: 5.7%; }}
.hourly-table th.th-day {{ width: 7.5%; }}

.hourly-table td {{
    padding: 7px 3px;
    color: #0f172a;
    text-align: center;
    border-bottom: 1px solid #cbd5e1;
    font-weight: 700;
    overflow: hidden;
    white-space: nowrap;
}}
.hourly-table td:first-child {{ 
    text-align: left; 
    padding-left: 10px; 
    white-space: normal;
    word-break: break-word;
    font-size: 12.5px;
    font-weight: 800;
}}
.hourly-cell {{
    font-family: 'JetBrains Mono', 'Consolas', monospace;
    font-size: 12.5px;
    letter-spacing: -0.3px;
}}
.hourly-val {{
    font-weight: 900;
    color: #0f172a;
    font-size: 13px;
}}
.cum-val {{
    color: #334155;
    font-size: 11px;
    font-weight: 700;
}}
.tr-tcf1-vin {{ background-color: #ffffff; }}
.tr-tcf1-vin:hover {{ background-color: #f0f9ff; }}
.tr-tcf2-vin {{ background-color: #ffffff; }}
.tr-tcf2-vin:hover {{ background-color: #f0fdf4; }}
.tr-tcf1-drop {{ background-color: #ffffff; }}
.tr-tcf1-drop:hover {{ background-color: #f0fdfa; }}
.tr-tcf2-drop {{ background-color: #ffffff; }}
.tr-tcf2-drop:hover {{ background-color: #fff7ed; }}

.tr-tot-vin td {{
    background: linear-gradient(90deg, #1e40af 0%, #3b82f6 100%) !important;
    color: #ffffff !important;
    font-weight: 900 !important;
    font-size: 13px !important;
    border-top: 2px solid #1d4ed8 !important;
    border-bottom: 2px solid #1d4ed8 !important;
}}
.tr-tot-vin .cum-val {{ color: #dbeafe !important; font-weight: 800 !important; }}
.tr-tot-vin .hourly-val {{ color: #ffffff !important; font-size: 13.5px !important; font-weight: 900 !important; }}

.tr-tot-drop td {{
    background: linear-gradient(90deg, #0f766e 0%, #14b8a6 100%) !important;
    color: #ffffff !important;
    font-weight: 900 !important;
    font-size: 13px !important;
    border-top: 2px solid #0d9488 !important;
    border-bottom: 2px solid #0d9488 !important;
}}
.tr-tot-drop .cum-val {{ color: #ccfbf1 !important; font-weight: 800 !important; }}
.tr-tot-drop .hourly-val {{ color: #ffffff !important; font-size: 13.5px !important; font-weight: 900 !important; }}

.badge-vin1 {{ background: #dbeafe; color: #1e40af; padding: 2px 7px; border-radius: 4px; font-size: 11px; font-weight: 800; display: inline-block; }}
.badge-vin2 {{ background: #e0e7ff; color: #4338ca; padding: 2px 7px; border-radius: 4px; font-size: 11px; font-weight: 800; display: inline-block; }}
.badge-drop1 {{ background: #dcfce7; color: #15803d; padding: 2px 7px; border-radius: 4px; font-size: 11px; font-weight: 800; display: inline-block; }}
.badge-drop2 {{ background: #ffedd5; color: #c2410c; padding: 2px 7px; border-radius: 4px; font-size: 11px; font-weight: 800; display: inline-block; }}
</style>
<div class="hourly-card">
<table class="hourly-table">
<thead>
<tr>"""
        for i, col_hdr in enumerate(all_headers):
            th_cls = "th-first" if i == 0 else ("th-day" if 'DAY' in str(col_hdr).upper() else ("th-shift" if 'SHIFT' in str(col_hdr).upper() else "th-time"))
            html_hourly_table += f"<th class='{th_cls}'>{_fmt_hdr_html(col_hdr)}</th>"
        html_hourly_table += "</tr></thead><tbody>"

        for r_type, r_label, r_class, r_data in table_rows_display:
            html_hourly_table += f'<tr class="tr-{r_class}">'
            if r_type == 'TCF1_VIN_GENERATION':
                html_hourly_table += f'<td><span class="badge-vin1">TCF1</span> &nbsp; <b>VIN Gen</b></td>'
            elif r_type == 'TCF2_VIN_GENERATION':
                html_hourly_table += f'<td><span class="badge-vin2">TCF2</span> &nbsp; <b>VIN Gen</b></td>'
            elif r_type == 'TCF1-DROP':
                html_hourly_table += f'<td><span class="badge-drop1">TCF1</span> &nbsp; <b>Drop</b></td>'
            elif r_type == 'TCF2-DROP':
                html_hourly_table += f'<td><span class="badge-drop2">TCF2</span> &nbsp; <b>Drop</b></td>'
            elif r_type == 'TOTAL_VIN':
                html_hourly_table += f'<td>🔹 <b>TOTAL VIN</b></td>'
            elif r_type == 'TOTAL_DROP':
                html_hourly_table += f'<td>🔸 <b>TOTAL DROP</b></td>'
            else:
                html_hourly_table += f'<td><b>{r_label}</b></td>'

            for c in time_cols:
                val_str = str(r_data.get(c, '0[0]'))
                h, cum = _parse_cell_val(val_str)
                html_hourly_table += f'<td class="hourly-cell"><span class="hourly-val">{h}</span><span class="cum-val"> [{cum}]</span></td>'

            for c in tot_cols:
                val = r_data.get(c, 0)
                html_hourly_table += f'<td><b>{val}</b></td>'

            html_hourly_table += '</tr>'

        html_hourly_table += "</tbody></table></div>"
        st.markdown(html_hourly_table, unsafe_allow_html=True)

        # Visual Analytics: TCF1 and TCF2 Attractive Separate Dropping Charts
        with st.expander("📈 View Hourly Dropping Line Charts (TCF1 & TCF2 Separate Lines)", expanded=False):
            tcf1_drop_data = []
            tcf2_drop_data = []
            for c in time_cols:
                h1_d, _ = _parse_cell_val(tcf1_drop[c]) if tcf1_drop is not None else (0, 0)
                h2_d, _ = _parse_cell_val(tcf2_drop[c]) if tcf2_drop is not None else (0, 0)
                m = re.match(r'(\d{1,2})(?::(\d{2}))?\s*([AP]M)?\s*-\s*(\d{1,2})(?::(\d{2}))?\s*([AP]M)', str(c), re.I)
                if m:
                    h1, m1, p1, h2, m2, p2 = m.groups()
                    t1 = f"{h1}:{m1}" if m1 and m1 != '00' else h1
                    t2 = f"{h2}:{m2}" if m2 and m2 != '00' else h2
                    period = p2 or p1 or ""
                    slot_label = f"{t1}-{t2} {period.upper()}"
                else:
                    slot_label = str(c).replace(':00', '').replace(' ', '')
                tcf1_drop_data.append({
                    'Time Slot': slot_label,
                    'Hourly Cabs Dropped': h1_d
                })
                tcf2_drop_data.append({
                    'Time Slot': slot_label,
                    'Hourly Cabs Dropped': h2_d
                })

            def _build_attractive_drop_chart(data_list, color_theme):
                df_c = pd.DataFrame(data_list)
                max_v = df_c['Hourly Cabs Dropped'].max() if not df_c.empty else 10
                y_max = max(max_v * 1.35, max_v + 6)
                
                base_c = alt.Chart(df_c).encode(
                    x=alt.X('Time Slot:N', sort=None, title=None, axis=alt.Axis(labelAngle=-25, labelFontSize=11, labelFontWeight='bold', labelColor='#334155')),
                    y=alt.Y('Hourly Cabs Dropped:Q', title='Cabs / Hour', scale=alt.Scale(domain=[0, y_max]), axis=alt.Axis(grid=True, gridColor='rgba(0,0,0,0.06)', labelFontSize=11, labelFontWeight='bold'))
                )
                
                area_c = base_c.mark_area(
                    opacity=0.18,
                    color=color_theme
                )
                
                line_c = base_c.mark_line(
                    color=color_theme,
                    size=3.5,
                    interpolate='monotone'
                )
                
                points_c = base_c.mark_circle(
                    size=120,
                    color=color_theme,
                    opacity=1
                )
                
                points_inner = base_c.mark_circle(
                    size=35,
                    color='#ffffff',
                    opacity=1
                )
                
                text_c = base_c.mark_text(
                    align='center',
                    baseline='bottom',
                    dy=-10,
                    fontSize=13,
                    fontWeight='bold',
                    color=color_theme
                ).encode(
                    text='Hourly Cabs Dropped:Q'
                )
                
                final_chart = (area_c + line_c + points_c + points_inner + text_c).properties(
                    height=270
                ).configure_view(
                    strokeWidth=0
                )
                return final_chart

            chart_col1, chart_col2 = st.columns(2)
            with chart_col1:
                st.markdown("##### 🟢 TCF1 Line Hourly Dropping Trend")
                if tcf1_drop_data:
                    chart_tcf1 = _build_attractive_drop_chart(tcf1_drop_data, '#059669')
                    st.altair_chart(chart_tcf1, use_container_width=True)
            with chart_col2:
                st.markdown("##### 🟠 TCF2 Line Hourly Dropping Trend")
                if tcf2_drop_data:
                    chart_tcf2 = _build_attractive_drop_chart(tcf2_drop_data, '#ea580c')
                    st.altair_chart(chart_tcf2, use_container_width=True)

        # Excel Export for Hourly Production Tracker
        export_hourly_rows = []
        for r_type, r_label, r_class, r_data in table_rows_display:
            row_dict = {'ACHIEVEMENT POINT': r_label}
            for c in time_cols:
                row_dict[c] = r_data.get(c, '0[0]')
            for c in tot_cols:
                row_dict[c] = r_data.get(c, 0)
            export_hourly_rows.append(row_dict)
            
        df_export_hourly = pd.DataFrame(export_hourly_rows)
        
        buf_hourly = io.BytesIO()
        with pd.ExcelWriter(buf_hourly, engine='openpyxl') as writer_hourly:
            df_export_hourly.to_excel(writer_hourly, index=False, sheet_name='Hourly Production')
            ws_h = writer_hourly.sheets['Hourly Production']
            
            f_hdr = Font(name='Calibri', size=11, bold=True, color='000000')
            fill_hdr = PatternFill(start_color='FCE4D6', end_color='FCE4D6', fill_type='solid')
            f_sub = Font(name='Calibri', size=11, bold=True, color='000000')
            fill_sub_v = PatternFill(start_color='BDD7EE', end_color='BDD7EE', fill_type='solid')
            fill_sub_d = PatternFill(start_color='C6EFCE', end_color='C6EFCE', fill_type='solid')
            f_norm = Font(name='Calibri', size=11, color='000000')
            b_thin = Border(
                left=Side(style='thin', color='BFBFBF'),
                right=Side(style='thin', color='BFBFBF'),
                top=Side(style='thin', color='BFBFBF'),
                bottom=Side(style='thin', color='BFBFBF')
            )
            
            ws_h.row_dimensions[1].height = 28
            for c_i in range(1, len(df_export_hourly.columns) + 1):
                c_cell = ws_h.cell(row=1, column=c_i)
                c_cell.font = f_hdr
                c_cell.fill = fill_hdr
                c_cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
                c_cell.border = b_thin
                
            for r_i in range(2, len(df_export_hourly) + 2):
                ws_h.row_dimensions[r_i].height = 22
                pt_val = str(ws_h.cell(row=r_i, column=1).value).strip().upper()
                is_tot_v = 'TOTAL VIN' in pt_val
                is_tot_d = 'TOTAL TCF DROP' in pt_val or 'TOTAL DROP' in pt_val
                
                for c_i in range(1, len(df_export_hourly.columns) + 1):
                    cell_obj = ws_h.cell(row=r_i, column=c_i)
                    cell_obj.border = b_thin
                    if c_i == 1:
                        cell_obj.alignment = Alignment(horizontal='left', vertical='center')
                    else:
                        cell_obj.alignment = Alignment(horizontal='center', vertical='center')
                        
                    if is_tot_v:
                        cell_obj.font = f_sub
                        cell_obj.fill = fill_sub_v
                    elif is_tot_d:
                        cell_obj.font = f_sub
                        cell_obj.fill = fill_sub_d
                    else:
                        cell_obj.font = f_norm
                        
            for col in ws_h.columns:
                m_len = max(len(str(cell.value or '')) for cell in col)
                c_let = openpyxl.utils.get_column_letter(col[0].column)
                ws_h.column_dimensions[c_let].width = max(m_len + 4, 14)

        st.download_button(
            label="📥 Export Hourly Production Tracker to Excel",
            data=buf_hourly.getvalue(),
            file_name="Hourly_Production_Tracker.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            key="export_hourly_prod_btn"
        )
        st.markdown("---")

    st.markdown("### 📊 Paint Shop Float Summary")
    st.markdown("""
        This report displays the paint shop buffer status by stage and model, matching the exact layout of the paint shop tracker sheet.
    """)
    
    if paint_summary_dict or (float_df is not None and not float_df.empty):
        # Map product names to internal TCF models
        def get_summary_product_to_model(prod_name):
            prod = str(prod_name).strip().upper()
            if 'HORNBILL' in prod:
                return 'PUNCH'
            elif 'NOVA' in prod:
                return 'PUNCH.EV'
            elif 'ETURNA' in prod:
                return 'HARRIER.EV'
            elif 'GRAVITAS' in prod:
                return 'SAFARI'
            elif 'Q5' in prod:
                return 'HARRIER'
            elif 'TAYRONA' in prod:
                return 'SAFARI.EV'
            return 'UNKNOWN'
            
        def get_row_paint_stage(row):
            return ae.get_detailed_paint_summary_stage(row)
                
        stages_list = [
            'PBS FLOAT', 
            'PBS TO POLISHING', 
            'POLISHING TO TOPCOAT', 
            'TOPCOAT TO WETSANDING G ROOFBLACK', 
            'TOPCOAT TO WETSANDING G FRESH', 
            'WETSANDING G TO SEALANT', 
            'PT ENTRY TO SEALANT', 
            'BIW LIFTING G TO PT', 
            'PT BYPASS'
        ]
        
        tcf1_models = ['PUNCH', 'PUNCH.EV']
        tcf2_models = ['HARRIER.EV', 'SAFARI', 'HARRIER', 'SAFARI.EV']
        
        rows = []
        
        def get_today_vin_count(vgl_df, model_name):
            if vgl_df is None or vgl_df.empty:
                return 0
            col = 'Model_Family' if 'Model_Family' in vgl_df.columns else ('Model' if 'Model' in vgl_df.columns else None)
            if not col:
                return 0
            sub = vgl_df[vgl_df[col] == model_name]
            if sub.empty and 'Model' in vgl_df.columns and col != 'Model':
                sub = vgl_df[vgl_df['Model'] == model_name]
            if 'VIN_Count' in sub.columns:
                return int(sub['VIN_Count'].sum())
            return len(sub)
            
        if paint_summary_dict:
            # Extract numbers directly from PPC_Float_Report_Paint_...
            # TCF1 Line
            for model in tcf1_models:
                m_dict = paint_summary_dict.get(model, {})
                today_vin = get_today_vin_count(tcf1_drops, model)
                row_data = {
                    'Paint Float': 'TCF1',
                    'MODEL': model,
                    'Today VIN': today_vin
                }
                for stage in stages_list:
                    row_data[stage] = m_dict.get(stage, 0)
                row_data['TOTAL UPTO SEALANT'] = m_dict.get('TOTAL UPTO SEALANT', (
                    row_data['PBS FLOAT'] + 
                    row_data['PBS TO POLISHING'] + 
                    row_data['POLISHING TO TOPCOAT'] + 
                    row_data['TOPCOAT TO WETSANDING G ROOFBLACK'] + 
                    row_data['TOPCOAT TO WETSANDING G FRESH'] + 
                    row_data['WETSANDING G TO SEALANT']
                ))
                row_data['TOTAL FLOAT'] = m_dict.get('TOTAL FLOAT', row_data['TOTAL UPTO SEALANT'] + row_data.get('PT ENTRY TO SEALANT', 0) + row_data.get('BIW LIFTING G TO PT', 0) + row_data.get('PT BYPASS', 0))
                rows.append(row_data)

            # TCF1 TOTAL
            tcf1_subtotal = {
                'Paint Float': 'TCF1',
                'MODEL': 'TCF1 TOTAL',
                'Today VIN': sum(r['Today VIN'] for r in rows if r['Paint Float'] == 'TCF1')
            }
            for col in ['TOTAL FLOAT'] + stages_list + ['TOTAL UPTO SEALANT']:
                tcf1_subtotal[col] = sum(r[col] for r in rows if r['Paint Float'] == 'TCF1')
            rows.append(tcf1_subtotal)

            # TCF2 Line
            tcf2_rows_start_idx = len(rows)
            for model in tcf2_models:
                m_dict = paint_summary_dict.get(model, {})
                today_vin = get_today_vin_count(tcf2_drops, model)
                row_data = {
                    'Paint Float': 'TCF2',
                    'MODEL': model,
                    'Today VIN': today_vin
                }
                for stage in stages_list:
                    row_data[stage] = m_dict.get(stage, 0)
                row_data['TOTAL UPTO SEALANT'] = m_dict.get('TOTAL UPTO SEALANT', (
                    row_data['PBS FLOAT'] + 
                    row_data['PBS TO POLISHING'] + 
                    row_data['POLISHING TO TOPCOAT'] + 
                    row_data['TOPCOAT TO WETSANDING G ROOFBLACK'] + 
                    row_data['TOPCOAT TO WETSANDING G FRESH'] + 
                    row_data['WETSANDING G TO SEALANT']
                ))
                row_data['TOTAL FLOAT'] = m_dict.get('TOTAL FLOAT', row_data['TOTAL UPTO SEALANT'] + row_data.get('PT ENTRY TO SEALANT', 0) + row_data.get('BIW LIFTING G TO PT', 0) + row_data.get('PT BYPASS', 0))
                rows.append(row_data)

            # TCF2 TOTAL
            tcf2_subtotal = {
                'Paint Float': 'TCF2',
                'MODEL': 'TCF2 TOTAL',
                'Today VIN': sum(r['Today VIN'] for r in rows[tcf2_rows_start_idx:] if r['Paint Float'] == 'TCF2')
            }
            for col in ['TOTAL FLOAT'] + stages_list + ['TOTAL UPTO SEALANT']:
                tcf2_subtotal[col] = sum(r[col] for r in rows[tcf2_rows_start_idx:] if r['Paint Float'] == 'TCF2')
            rows.append(tcf2_subtotal)

            # GRAND TOTAL
            grand_total = {
                'Paint Float': '',
                'MODEL': 'GRAND TOTAL',
                'Today VIN': tcf1_subtotal['Today VIN'] + tcf2_subtotal['Today VIN']
            }
            for col in ['TOTAL FLOAT'] + stages_list + ['TOTAL UPTO SEALANT']:
                grand_total[col] = tcf1_subtotal[col] + tcf2_subtotal[col]
            rows.append(grand_total)

            summary_df = pd.DataFrame(rows)
        elif float_df is not None and not float_df.empty:
            temp_float_df = float_df.copy()
            temp_float_df['Model_Mapped'] = temp_float_df['PRODUCT'].apply(get_summary_product_to_model)
            temp_float_df['Stage'] = temp_float_df.apply(get_row_paint_stage, axis=1)
            
            # TCF1 Line
            tcf1_sub_df = temp_float_df[temp_float_df['SHOP'] == 'TCF1']
            for model in tcf1_models:
                model_df = tcf1_sub_df[tcf1_sub_df['Model_Mapped'] == model]
                today_vin = get_today_vin_count(tcf1_drops, model)
                
                row_data = {
                    'Paint Float': 'TCF1',
                    'MODEL': model,
                    'Today VIN': today_vin
                }
                
                total_float = 0
                for stage in stages_list:
                    cnt = len(model_df[model_df['Stage'] == stage])
                    row_data[stage] = cnt
                    total_float += cnt
                    
                row_data['TOTAL FLOAT'] = total_float
                row_data['TOTAL UPTO SEALANT'] = (
                    row_data['PBS FLOAT'] + 
                    row_data['PBS TO POLISHING'] + 
                    row_data['POLISHING TO TOPCOAT'] + 
                    row_data['TOPCOAT TO WETSANDING G ROOFBLACK'] + 
                    row_data['TOPCOAT TO WETSANDING G FRESH'] + 
                    row_data['WETSANDING G TO SEALANT']
                )
                rows.append(row_data)
                
            # TCF1 TOTAL
            tcf1_subtotal = {
                'Paint Float': 'TCF1',
                'MODEL': 'TCF1 TOTAL',
                'Today VIN': sum(r['Today VIN'] for r in rows if r['Paint Float'] == 'TCF1')
            }
            for col in ['TOTAL FLOAT'] + stages_list + ['TOTAL UPTO SEALANT']:
                tcf1_subtotal[col] = sum(r[col] for r in rows if r['Paint Float'] == 'TCF1')
            rows.append(tcf1_subtotal)
            
            # TCF2 Line
            tcf2_sub_df = temp_float_df[temp_float_df['SHOP'] == 'TCF2']
            tcf2_rows_start_idx = len(rows)
            for model in tcf2_models:
                model_df = tcf2_sub_df[tcf2_sub_df['Model_Mapped'] == model]
                today_vin = get_today_vin_count(tcf2_drops, model)
                
                row_data = {
                    'Paint Float': 'TCF2',
                    'MODEL': model,
                    'Today VIN': today_vin
                }
                
                total_float = 0
                for stage in stages_list:
                    cnt = len(model_df[model_df['Stage'] == stage])
                    row_data[stage] = cnt
                    total_float += cnt
                    
                row_data['TOTAL FLOAT'] = total_float
                row_data['TOTAL UPTO SEALANT'] = (
                    row_data['PBS FLOAT'] + 
                    row_data['PBS TO POLISHING'] + 
                    row_data['POLISHING TO TOPCOAT'] + 
                    row_data['TOPCOAT TO WETSANDING G ROOFBLACK'] + 
                    row_data['TOPCOAT TO WETSANDING G FRESH'] + 
                    row_data['WETSANDING G TO SEALANT']
                )
                rows.append(row_data)
                
            # TCF2 TOTAL
            tcf2_subtotal = {
                'Paint Float': 'TCF2',
                'MODEL': 'TCF2 TOTAL',
                'Today VIN': sum(r['Today VIN'] for r in rows[tcf2_rows_start_idx:] if r['Paint Float'] == 'TCF2')
            }
            for col in ['TOTAL FLOAT'] + stages_list + ['TOTAL UPTO SEALANT']:
                tcf2_subtotal[col] = sum(r[col] for r in rows[tcf2_rows_start_idx:] if r['Paint Float'] == 'TCF2')
            rows.append(tcf2_subtotal)
            
            # GRAND TOTAL
            grand_total = {
                'Paint Float': '',
                'MODEL': 'GRAND TOTAL',
                'Today VIN': tcf1_subtotal['Today VIN'] + tcf2_subtotal['Today VIN']
            }
            for col in ['TOTAL FLOAT'] + stages_list + ['TOTAL UPTO SEALANT']:
                grand_total[col] = tcf1_subtotal[col] + tcf2_subtotal[col]
            rows.append(grand_total)
            
            summary_df = pd.DataFrame(rows)
        else:
            summary_df = pd.DataFrame()
        
        # Rename columns to match user copy perfectly
        display_col_mapping = {
            'Paint Float': 'Paint Float',
            'MODEL': 'MODEL',
            'TOTAL FLOAT': 'TOTAL FLOAT',
            'PBS FLOAT': 'PBS FLOAT',
            'PBS TO POLISHING': 'PBS TO POLISHING',
            'POLISHING TO TOPCOAT': 'POLISHING TO TOPCOAT',
            'TOPCOAT TO WETSANDING G ROOFBLACK': 'TOPCOAT TO WETSANDING G ROOFBLACK',
            'TOPCOAT TO WETSANDING G FRESH': 'TOPCOAT TO WETSANDING G FRESH',
            'WETSANDING G TO SEALANT': 'WETSANDING G TO SEALANT',
            'TOTAL UPTO SEALANT': 'TOTAL UPTO SEALANT',
            'PT ENTRY TO SEALANT': 'PT ENTRY TO SEALANT',
            'BIW LIFTING G TO PT': 'BIW LIFTING G TO PT',
            'PT BYPASS': 'PT BYPASS',
            'Today VIN': 'Today VIN'
        }
        
        summary_df = summary_df[[
            'Paint Float', 'MODEL', 'TOTAL FLOAT', 'PBS FLOAT', 'PBS TO POLISHING',
            'POLISHING TO TOPCOAT', 'TOPCOAT TO WETSANDING G ROOFBLACK',
            'TOPCOAT TO WETSANDING G FRESH', 'WETSANDING G TO SEALANT',
            'TOTAL UPTO SEALANT', 'PT ENTRY TO SEALANT', 'BIW LIFTING G TO PT',
            'PT BYPASS', 'Today VIN'
        ]].rename(columns=display_col_mapping)
        
        # Helper to generate beautiful wrapped HTML table for summary float report
        is_dark_theme = st.session_state.get('theme', '☀️ White Theme') == '🌙 Dark Theme'
        
        def render_html_float_summary(df, is_dark):
            th_bg = "#1F2937" if is_dark else "#F3F4F6"
            th_text = "#FAFAFA" if is_dark else "#374151"
            td_border = "#30363D" if is_dark else "#E5E7EB"
            text_color = "#FAFAFA" if is_dark else "#111827"
            
            html = f"""
            <div style="overflow-x: auto; border: 1px solid {td_border}; border-radius: 12px; margin-bottom: 2rem; background-color: {'#161B22' if is_dark else '#FFFFFF'}; box-shadow: 0 4px 12px rgba(0,0,0,0.03);">
            <table style="width: 100%; border-collapse: collapse; font-family: 'Inter', sans-serif; font-size: 12px; color: {text_color};">
                <thead>
                    <tr style="background-color: {th_bg}; border-bottom: 2px solid {td_border};">
                        <th style="padding: 10px 8px; border: 1px solid {td_border}; text-align: left; color: {th_text}; font-weight: 600;">Paint Float</th>
                        <th style="padding: 10px 8px; border: 1px solid {td_border}; text-align: left; color: {th_text}; font-weight: 600; width: 110px;">MODEL</th>
                        <th style="padding: 10px 8px; border: 1px solid {td_border}; text-align: center; color: {th_text}; font-weight: 600; min-width: 60px; word-wrap: break-word; white-space: normal;">TOTAL FLOAT</th>
                        <th style="padding: 10px 8px; border: 1px solid {td_border}; text-align: center; color: {th_text}; font-weight: 600; min-width: 60px; word-wrap: break-word; white-space: normal;">PBS FLOAT</th>
                        <th style="padding: 10px 8px; border: 1px solid {td_border}; text-align: center; color: {th_text}; font-weight: 600; min-width: 80px; word-wrap: break-word; white-space: normal;">PBS TO POLISHING</th>
                        <th style="padding: 10px 8px; border: 1px solid {td_border}; text-align: center; color: {th_text}; font-weight: 600; min-width: 80px; word-wrap: break-word; white-space: normal;">POLISHING TO TOPCOAT</th>
                        <th style="padding: 10px 8px; border: 1px solid {td_border}; text-align: center; color: {th_text}; font-weight: 600; min-width: 100px; max-width: 120px; word-wrap: break-word; white-space: normal;">TOPCOAT TO WETSANDING G ROOFBLACK</th>
                        <th style="padding: 10px 8px; border: 1px solid {td_border}; text-align: center; color: {th_text}; font-weight: 600; min-width: 100px; max-width: 120px; word-wrap: break-word; white-space: normal;">TOPCOAT TO WETSANDING G FRESH</th>
                        <th style="padding: 10px 8px; border: 1px solid {td_border}; text-align: center; color: {th_text}; font-weight: 600; min-width: 100px; max-width: 120px; word-wrap: break-word; white-space: normal;">WETSANDING G TO SEALANT</th>
                        <th style="padding: 10px 8px; border: 1px solid {td_border}; text-align: center; color: {th_text}; font-weight: 600; min-width: 80px; max-width: 100px; word-wrap: break-word; white-space: normal;">TOTAL UPTO SEALANT</th>
                        <th style="padding: 10px 8px; border: 1px solid {td_border}; text-align: center; color: {th_text}; font-weight: 600; min-width: 80px; max-width: 100px; word-wrap: break-word; white-space: normal;">PT ENTRY TO SEALANT</th>
                        <th style="padding: 10px 8px; border: 1px solid {td_border}; text-align: center; color: {th_text}; font-weight: 600; min-width: 80px; max-width: 100px; word-wrap: break-word; white-space: normal;">BIW LIFTING G TO PT</th>
                        <th style="padding: 10px 8px; border: 1px solid {td_border}; text-align: center; color: {th_text}; font-weight: 600; min-width: 60px; word-wrap: break-word; white-space: normal;">PT BYPASS</th>
                        <th style="padding: 10px 8px; border: 1px solid {td_border}; text-align: center; color: {th_text}; font-weight: 600; min-width: 60px; word-wrap: break-word; white-space: normal;">Today VIN</th>
                    </tr>
                </thead>
                <tbody>
            """
            
            for idx_r, row_r in df.iterrows():
                model_val = str(row_r.get('MODEL', '')).strip()
                
                row_bg = "transparent"
                row_text = text_color
                font_weight = "normal"
                
                if 'TOTAL' in model_val and 'GRAND' not in model_val:
                    row_bg = "#3b1f3c" if is_dark else "#f2dcdb"
                    row_text = "#f2dcdb" if is_dark else "#5c1d1b"
                    font_weight = "bold"
                elif 'GRAND TOTAL' in model_val:
                    row_bg = "#4a3f00" if is_dark else "#ffffc5"
                    row_text = "#ffff00" if is_dark else "#806000"
                    font_weight = "bold"
                    
                html += f'<tr style="background-color: {row_bg}; color: {row_text}; font-weight: {font_weight}; border-bottom: 1px solid {td_border};">'
                html += f'<td style="padding: 8px; border: 1px solid {td_border}; text-align: left;">{row_r.get("Paint Float", "")}</td>'
                html += f'<td style="padding: 8px; border: 1px solid {td_border}; text-align: left;">{row_r.get("MODEL", "")}</td>'
                
                for col in ['TOTAL FLOAT', 'PBS FLOAT', 'PBS TO POLISHING', 'POLISHING TO TOPCOAT', 
                            'TOPCOAT TO WETSANDING G ROOFBLACK', 'TOPCOAT TO WETSANDING G FRESH', 
                            'WETSANDING G TO SEALANT', 'TOTAL UPTO SEALANT', 'PT ENTRY TO SEALANT', 
                            'BIW LIFTING G TO PT', 'PT BYPASS', 'Today VIN']:
                    val = row_r.get(col, 0)
                    val_str = str(val) if pd.notna(val) else "0"
                    html += f'<td style="padding: 8px; border: 1px solid {td_border}; text-align: center;">{val_str}</td>'
                html += '</tr>'
                
            html += """
                </tbody>
            </table>
            </div>
            """
            return html
            
        # Display the first table (wrapped nicely in HTML)
        st.markdown(render_html_float_summary(summary_df, is_dark_theme), unsafe_allow_html=True)
        
        # ----------------- ENGINE & BATTERY REQUIREMENT SUMMARY REPORT -----------------
        st.markdown("<div style='height: 1.5rem;'></div>", unsafe_allow_html=True)
        st.markdown("### 📊 Engine & Battery Requirement Summary")
        st.markdown("""
            This report summarizes raw engine inventory status against paint shop float and computes clear-to-build requirements.
        """)
        
        # Compute dictionary values for engines
        engine_stocks_dict = {}
        engine_ta_dict = {}
        if 'engine_df' in st.session_state and st.session_state.engine_df is not None:
            for idx, r_eng in st.session_state.engine_df.iterrows():
                p_no = str(r_eng['Engine Part No']).strip()
                engine_stocks_dict[p_no] = int(r_eng.get('Clearance After 6:30AM', 0))
                engine_ta_dict[p_no] = str(r_eng.get('TA Code', '—')).strip()
                
        # Create mapping dictionary from Short Vehicle Code -> Engine
        vc_to_engine = {}
        if bom_df is not None and not bom_df.empty:
            vc_to_engine = dict(zip(bom_df['Short Vehicle Code'].astype(str).str.strip(), bom_df['Engine'].astype(str).str.strip()))
            
        # Today VIN per engine part
        today_vin_dict = {}
        for vgl_df in [tcf1_drops, tcf2_drops]:
            if vgl_df is not None and not vgl_df.empty:
                vc_col = 'VEHICLE CODE' if 'VEHICLE CODE' in vgl_df.columns else ('VC' if 'VC' in vgl_df.columns else None)
                if vc_col:
                    vgl_df['Engine_Part'] = vgl_df[vc_col].astype(str).str.strip().str[:9].map(vc_to_engine)
                    for part in vgl_df['Engine_Part'].dropna().unique():
                        p_str = str(part).strip()
                        if p_str in ['None', 'nan', '0', '']:
                            continue
                        sub = vgl_df[vgl_df['Engine_Part'] == part]
                        cnt = int(sub['VIN_Count'].sum()) if 'VIN_Count' in sub.columns else len(sub)
                        today_vin_dict[p_str] = today_vin_dict.get(p_str, 0) + cnt
                
        # Float demands per engine part
        pbs_float_dict = {}
        upto_sealant_dict = {}
        total_float_dict = {}
        
        stages_upto_sealant = [
            'PBS FLOAT', 
            'PBS TO POLISHING', 
            'POLISHING TO TOPCOAT', 
            'TOPCOAT TO WETSANDING G ROOFBLACK', 
            'TOPCOAT TO WETSANDING G FRESH', 
            'WETSANDING G TO SEALANT'
        ]
        
        if paint_summary_vc_dict:
            for svc, counts in paint_summary_vc_dict.items():
                p_str = vc_to_engine.get(svc)
                if p_str and str(p_str).strip() not in ['None', 'nan', '0']:
                    p_clean = str(p_str).strip()
                    total_float_dict[p_clean] = total_float_dict.get(p_clean, 0) + counts['TOTAL FLOAT']
                    pbs_float_dict[p_clean] = pbs_float_dict.get(p_clean, 0) + counts['PBS FLOAT']
                    upto_sealant_dict[p_clean] = upto_sealant_dict.get(p_clean, 0) + counts['TOTAL UPTO SEALANT']
        elif temp_float_df is not None and not temp_float_df.empty:
            vc_col = 'VEHICLE CODE' if 'VEHICLE CODE' in temp_float_df.columns else ('VC' if 'VC' in temp_float_df.columns else None)
            if vc_col:
                temp_float_df['Engine_Part'] = temp_float_df[vc_col].astype(str).str.strip().str[:9].map(vc_to_engine)
                for idx, row_f in temp_float_df.iterrows():
                    part = row_f.get('Engine_Part')
                    if pd.isna(part):
                        continue
                    p_str = str(part).strip()
                    stage = row_f.get('Stage', '')
                    
                    total_float_dict[p_str] = total_float_dict.get(p_str, 0) + 1
                    if stage == 'PBS FLOAT':
                        pbs_float_dict[p_str] = pbs_float_dict.get(p_str, 0) + 1
                    if stage in stages_upto_sealant:
                        upto_sealant_dict[p_str] = upto_sealant_dict.get(p_str, 0) + 1
                
        # Build TCF1 rows
        punch_parts = [
            ("54850000PTP001", "Punch MT SA"),
            ("54850000PTP002", "Punch AMT SA"),
            ("54970000PTP002", "Punch TC MCE"),
            ("54970000PTP003", "Punch MCE MT"),
            ("54970000PTP004", "Punch MCE AMT"),
            ("54970000PTP005", "Punch MCE CNG MT"),
            ("54970000PTP031", "Punch MCE CNG AMT")
        ]
        
        table2_rows = []
        for part, model in punch_parts:
            clearance = engine_stocks_dict.get(part, 0)
            today_vin = today_vin_dict.get(part, 0)
            bal = clearance - today_vin
            pbs = pbs_float_dict.get(part, 0)
            sealant = upto_sealant_dict.get(part, 0)
            total = total_float_dict.get(part, 0)
            table2_rows.append({
                'Engine Part No': part,
                'Model': model,
                'TA Code': engine_ta_dict.get(part, '—'),
                'Clearance After 6:30AM': clearance,
                'Today VIN': today_vin,
                'Bal': bal,
                'PBS FLOAT': pbs,
                'Float UPTO SEALANT': sealant,
                'TOTAL FLOAT': total,
                'With respect to PBS FLOAT': bal - pbs,
                'With respect to Sealant FLOAT': bal - sealant,
                'With respect to Total FLOAT': bal - total,
                'Type': 'row'
            })
            
        subtotal_1_2 = {
            'Engine Part No': '',
            'Model': '1.2 Lit Total',
            'TA Code': '',
            'Clearance After 6:30AM': '',
            'Today VIN': sum(r['Today VIN'] for r in table2_rows),
            'Bal': '',
            'PBS FLOAT': sum(r['PBS FLOAT'] for r in table2_rows),
            'Float UPTO SEALANT': sum(r['Float UPTO SEALANT'] for r in table2_rows),
            'TOTAL FLOAT': sum(r['TOTAL FLOAT'] for r in table2_rows),
            'With respect to PBS FLOAT': '',
            'With respect to Sealant FLOAT': '',
            'With respect to Total FLOAT': '',
            'Type': 'subtotal'
        }
        table2_rows.append(subtotal_1_2)
        
        # Nova
        part_nova = "546816111212"
        model_nova = "Nova"
        if 'nova_materials_df' in st.session_state and st.session_state.nova_materials_df is not None and not st.session_state.nova_materials_df.empty:
            clearance_nova = int(st.session_state.nova_materials_df['Clearance Qty'].min())
        else:
            clearance_nova = 182
        today_vin_nova = today_vin_dict.get(part_nova, 0)
        bal_nova = clearance_nova - today_vin_nova
        pbs_nova = pbs_float_dict.get(part_nova, 0)
        sealant_nova = upto_sealant_dict.get(part_nova, 0)
        total_nova = total_float_dict.get(part_nova, 0)
        row_nova = {
            'Engine Part No': part_nova,
            'Model': model_nova,
            'TA Code': engine_ta_dict.get(part_nova, '5468'),
            'Clearance After 6:30AM': clearance_nova,
            'Today VIN': today_vin_nova,
            'Bal': bal_nova,
            'PBS FLOAT': pbs_nova,
            'Float UPTO SEALANT': sealant_nova,
            'TOTAL FLOAT': total_nova,
            'With respect to PBS FLOAT': bal_nova - pbs_nova,
            'With respect to Sealant FLOAT': bal_nova - sealant_nova,
            'With respect to Total FLOAT': bal_nova - total_nova,
            'Type': 'row'
        }
        table2_rows.append(row_nova)

        # TCF1 Grand Total
        tcf1_grand = {
            'Engine Part No': '',
            'Model': 'TCF1',
            'TA Code': '',
            'Clearance After 6:30AM': '',
            'Today VIN': subtotal_1_2['Today VIN'] + row_nova['Today VIN'],
            'Bal': '',
            'PBS FLOAT': subtotal_1_2['PBS FLOAT'] + row_nova['PBS FLOAT'],
            'Float UPTO SEALANT': subtotal_1_2['Float UPTO SEALANT'] + row_nova['Float UPTO SEALANT'],
            'TOTAL FLOAT': subtotal_1_2['TOTAL FLOAT'] + row_nova['TOTAL FLOAT'],
            'With respect to PBS FLOAT': '',
            'With respect to Sealant FLOAT': '',
            'With respect to Total FLOAT': '',
            'Type': 'total'
        }
        table2_rows.append(tcf1_grand)
        
        # Build TCF2 rows
        tcf2_parts = [
            ("572900000118", "Harrier / Safari Diesel AT"),
            ("572900000120", "Harrier / Safari Diesel MT"),
            ("54780000PTP001", "Harrier / Safari Petrol TGDI MT"),
            ("54780000PTP002", "Harrier / Safari Petrol TGDI AT")
        ]
        
        tcf2_start_idx = len(table2_rows)
        for part, model in tcf2_parts:
            clearance = engine_stocks_dict.get(part, 0)
            today_vin = today_vin_dict.get(part, 0)
            bal = clearance - today_vin
            pbs = pbs_float_dict.get(part, 0)
            sealant = upto_sealant_dict.get(part, 0)
            total = total_float_dict.get(part, 0)
            table2_rows.append({
                'Engine Part No': part,
                'Model': model,
                'TA Code': engine_ta_dict.get(part, '—'),
                'Clearance After 6:30AM': clearance,
                'Today VIN': today_vin,
                'Bal': bal,
                'PBS FLOAT': pbs,
                'Float UPTO SEALANT': sealant,
                'TOTAL FLOAT': total,
                'With respect to PBS FLOAT': bal - pbs,
                'With respect to Sealant FLOAT': bal - sealant,
                'With respect to Total FLOAT': bal - total,
                'Type': 'row'
            })
            
        subtotal_2_0 = {
            'Engine Part No': '',
            'Model': '2 Lit Total',
            'TA Code': '',
            'Clearance After 6:30AM': '',
            'Today VIN': sum(r['Today VIN'] for r in table2_rows[tcf2_start_idx:]),
            'Bal': '',
            'PBS FLOAT': sum(r['PBS FLOAT'] for r in table2_rows[tcf2_start_idx:]),
            'Float UPTO SEALANT': sum(r['Float UPTO SEALANT'] for r in table2_rows[tcf2_start_idx:]),
            'TOTAL FLOAT': sum(r['TOTAL FLOAT'] for r in table2_rows[tcf2_start_idx:]),
            'With respect to PBS FLOAT': '',
            'With respect to Sealant FLOAT': '',
            'With respect to Total FLOAT': '',
            'Type': 'subtotal'
        }
        table2_rows.append(subtotal_2_0)
        
        # Harrier EV
        part_hev = "547380400103"
        model_hev = "Harrier EV"
        clearance_hev = 160
        today_vin_hev = today_vin_dict.get(part_hev, 0)
        bal_hev = clearance_hev - today_vin_hev
        pbs_hev = pbs_float_dict.get(part_hev, 0)
        sealant_hev = upto_sealant_dict.get(part_hev, 0)
        total_hev = total_float_dict.get(part_hev, 0)
        row_hev = {
            'Engine Part No': part_hev,
            'Model': model_hev,
            'TA Code': engine_ta_dict.get(part_hev, '5473'),
            'Clearance After 6:30AM': clearance_hev,
            'Today VIN': today_vin_hev,
            'Bal': bal_hev,
            'PBS FLOAT': pbs_hev,
            'Float UPTO SEALANT': sealant_hev,
            'TOTAL FLOAT': total_hev,
            'With respect to PBS FLOAT': bal_hev - pbs_hev,
            'With respect to Sealant FLOAT': bal_hev - sealant_hev,
            'With respect to Total FLOAT': bal_hev - total_hev,
            'Type': 'row'
        }
        table2_rows.append(row_hev)
        
        # TCF2 Grand Total
        tcf2_grand = {
            'Engine Part No': '',
            'Model': 'TCF2',
            'TA Code': '',
            'Clearance After 6:30AM': '',
            'Today VIN': subtotal_2_0['Today VIN'] + row_hev['Today VIN'],
            'Bal': '',
            'PBS FLOAT': subtotal_2_0['PBS FLOAT'] + row_hev['PBS FLOAT'],
            'Float UPTO SEALANT': subtotal_2_0['Float UPTO SEALANT'] + row_hev['Float UPTO SEALANT'],
            'TOTAL FLOAT': subtotal_2_0['TOTAL FLOAT'] + row_hev['TOTAL FLOAT'],
            'With respect to PBS FLOAT': '',
            'With respect to Sealant FLOAT': '',
            'With respect to Total FLOAT': '',
            'Type': 'total'
        }
        table2_rows.append(tcf2_grand)
        
        # Render Table 2 in beautiful HTML with rowspan/colspan
        def render_html_table_2(rows, is_dark):
            th_bg = "#1F2937" if is_dark else "#F3F4F6"
            th_text = "#FAFAFA" if is_dark else "#374151"
            td_border = "#30363D" if is_dark else "#E5E7EB"
            text_color = "#FAFAFA" if is_dark else "#111827"
            
            clearance_bg = "#1b4d32" if is_dark else "#d8f3e5"
            clearance_text = "#FAFAFA" if is_dark else "#1b4d32"
            
            bal_bg = "#4a274c" if is_dark else "#f2dcdb"
            bal_text = "#FAFAFA" if is_dark else "#5c1d1b"
            
            alert_bg = "#5c1d1d" if is_dark else "#ffd1d1"
            alert_text = "#FAFAFA" if is_dark else "#5c1d1d"
            
            html = f"""
            <div style="overflow-x: auto; border: 1px solid {td_border}; border-radius: 12px; margin-bottom: 2rem; background-color: {'#161B22' if is_dark else '#FFFFFF'}; box-shadow: 0 4px 12px rgba(0,0,0,0.03);">
            <table style="width: 100%; border-collapse: collapse; font-family: 'Inter', sans-serif; font-size: 12px; color: {text_color};">
                <thead>
                    <tr style="background-color: {th_bg}; border-bottom: 1px solid {td_border};">
                        <th rowspan="2" style="padding: 10px 8px; border: 1px solid {td_border}; text-align: center; color: {th_text}; font-weight: 600; vertical-align: middle;">Engine / Battery Part No</th>
                        <th rowspan="2" style="padding: 10px 8px; border: 1px solid {td_border}; text-align: left; color: {th_text}; font-weight: 600; width: 180px; vertical-align: middle;">Model</th>
                        <th rowspan="2" style="padding: 10px 8px; border: 1px solid {td_border}; text-align: center; color: {th_text}; font-weight: 600; vertical-align: middle;">TA Code</th>
                        <th rowspan="2" style="padding: 10px 8px; border: 1px solid {td_border}; text-align: center; color: {th_text}; font-weight: 600; min-width: 80px; white-space: normal; vertical-align: middle;">Clearance After 6:30AM</th>
                        <th rowspan="2" style="padding: 10px 8px; border: 1px solid {td_border}; text-align: center; color: {th_text}; font-weight: 600; vertical-align: middle;">Today VIN</th>
                        <th rowspan="2" style="padding: 10px 8px; border: 1px solid {td_border}; text-align: center; color: {th_text}; font-weight: 600; background-color: {bal_bg}; color: {bal_text}; vertical-align: middle;">Bal</th>
                        <th colspan="3" style="padding: 6px; border: 1px solid {td_border}; text-align: center; color: {th_text}; font-weight: 600;">Paint Float</th>
                        <th colspan="3" style="padding: 6px; border: 1px solid {td_border}; text-align: center; color: {th_text}; font-weight: 600;">Engine & Battery requirement</th>
                    </tr>
                    <tr style="background-color: {th_bg}; border-bottom: 2px solid {td_border};">
                        <th style="padding: 6px; border: 1px solid {td_border}; text-align: center; color: {th_text}; font-weight: 600;">PBS FLOAT</th>
                        <th style="padding: 6px; border: 1px solid {td_border}; text-align: center; color: {th_text}; font-weight: 600;">Float UPTO SEALANT</th>
                        <th style="padding: 6px; border: 1px solid {td_border}; text-align: center; color: {th_text}; font-weight: 600;">TOTAL FLOAT</th>
                        <th style="padding: 6px; border: 1px solid {td_border}; text-align: center; color: {th_text}; font-weight: 600; min-width: 90px; white-space: normal;">With respect to PBS FLOAT</th>
                        <th style="padding: 6px; border: 1px solid {td_border}; text-align: center; color: {th_text}; font-weight: 600; min-width: 90px; white-space: normal;">With respect to Sealant FLOAT</th>
                        <th style="padding: 6px; border: 1px solid {td_border}; text-align: center; color: {th_text}; font-weight: 600; min-width: 90px; white-space: normal;">With respect to Total FLOAT</th>
                    </tr>
                </thead>
                <tbody>
            """
            
            for r_data in rows:
                r_type = r_data['Type']
                
                row_bg = "transparent"
                row_text = text_color
                font_weight = "normal"
                
                if r_type == 'subtotal':
                    row_bg = "#005b8a" if is_dark else "#00B0F0"
                    row_text = "#FFFFFF"
                    font_weight = "bold"
                elif r_type == 'total':
                    row_bg = "#7f7f00" if is_dark else "#ffff00"
                    row_text = "#FAFAFA" if is_dark else "#000000"
                    font_weight = "bold"
                    
                html += f'<tr style="background-color: {row_bg}; color: {row_text}; font-weight: {font_weight}; border-bottom: 1px solid {td_border};">'
                html += f'<td style="padding: 8px; border: 1px solid {td_border}; text-align: center;">{r_data["Engine Part No"]}</td>'
                html += f'<td style="padding: 8px; border: 1px solid {td_border}; text-align: left;">{r_data["Model"]}</td>'
                html += f'<td style="padding: 8px; border: 1px solid {td_border}; text-align: center;">{r_data["TA Code"]}</td>'
                
                val_clearance = r_data["Clearance After 6:30AM"]
                if val_clearance != "" and r_type == 'row':
                    html += f'<td style="padding: 8px; border: 1px solid {td_border}; text-align: center; background-color: {clearance_bg}; color: {clearance_text}; font-weight: bold;">{val_clearance}</td>'
                else:
                    html += f'<td style="padding: 8px; border: 1px solid {td_border}; text-align: center;">{val_clearance}</td>'
                    
                html += f'<td style="padding: 8px; border: 1px solid {td_border}; text-align: center;">{r_data["Today VIN"]}</td>'
                
                val_bal = r_data["Bal"]
                if val_bal != "" and r_type == 'row':
                    html += f'<td style="padding: 8px; border: 1px solid {td_border}; text-align: center; background-color: {bal_bg}; color: {bal_text}; font-weight: bold;">{val_bal}</td>'
                else:
                    html += f'<td style="padding: 8px; border: 1px solid {td_border}; text-align: center;">{val_bal}</td>'
                    
                html += f'<td style="padding: 8px; border: 1px solid {td_border}; text-align: center;">{r_data["PBS FLOAT"]}</td>'
                html += f'<td style="padding: 8px; border: 1px solid {td_border}; text-align: center;">{r_data["Float UPTO SEALANT"]}</td>'
                html += f'<td style="padding: 8px; border: 1px solid {td_border}; text-align: center;">{r_data["TOTAL FLOAT"]}</td>'
                
                for col_k in ["With respect to PBS FLOAT", "With respect to Sealant FLOAT", "With respect to Total FLOAT"]:
                    val_req = r_data[col_k]
                    if val_req != "" and r_type == 'row':
                        if isinstance(val_req, (int, float)) and val_req < 0:
                            html += f'<td style="padding: 8px; border: 1px solid {td_border}; text-align: center; background-color: {alert_bg}; color: {alert_text}; font-weight: bold;">{val_req}</td>'
                        else:
                            html += f'<td style="padding: 8px; border: 1px solid {td_border}; text-align: center;">{val_req}</td>'
                    else:
                        html += f'<td style="padding: 8px; border: 1px solid {td_border}; text-align: center;">{val_req}</td>'
                        
                html += '</tr>'
            html += """
                </tbody>
            </table>
            </div>
            """
            return html
            
        st.markdown(render_html_table_2(table2_rows, is_dark_theme), unsafe_allow_html=True)
        
        # ----------------- COCKPIT & WIRING SHORTAGE DATA PREPARATION -----------------
        
        # Helper to build shortage table for Cockpit or Front Wiring
        def build_formatted_shortage_table(part_col_name, stock_tcf1, stock_tcf2, bom_df, float_df, paint_summary_vc_dict, tcf1_drops, tcf2_drops, only_shortage=True):
            if bom_df is None or bom_df.empty:
                return pd.DataFrame()
                
            local_engine_to_model = {}
            local_engine_to_line = {}
            if 'engine_df' in st.session_state and not st.session_state.engine_df.empty:
                local_engine_to_model = dict(zip(st.session_state.engine_df['Engine Part No'].astype(str).str.strip(), st.session_state.engine_df['Model']))
                local_engine_to_line = dict(zip(st.session_state.engine_df['Engine Part No'].astype(str).str.strip(), st.session_state.engine_df['TCF Line']))
            else:
                local_engine_to_model = {item['Engine Part No']: item['Model'] for item in engine_default_data}
                local_engine_to_line = {item['Engine Part No']: item['TCF Line'] for item in engine_default_data}

            local_engine_to_model['546816111212'] = 'Punch EV (Nova)'
            local_engine_to_line['546816111212'] = 'TCF1'
            local_engine_to_model['547380400103'] = 'Harrier EV'
            local_engine_to_line['547380400103'] = 'TCF2'

            vc_to_part = dict(zip(bom_df['Short Vehicle Code'].astype(str).str.strip(), bom_df[part_col_name].astype(str).str.strip()))
            
            part_to_models = {}
            part_to_vcs = {}
            part_to_line = {}
            
            for idx, row in bom_df.iterrows():
                eng = str(row.get('Engine', '')).strip()
                part = str(row.get(part_col_name, '')).strip()
                vc = str(row.get('Short Vehicle Code', '')).strip()
                mdl = local_engine_to_model.get(eng, '')
                line = local_engine_to_line.get(eng, '')
                if part and part not in ['None', 'nan', '0']:
                    if mdl:
                        part_to_models.setdefault(part, set()).add(mdl)
                    if vc and vc not in ['None', 'nan', '0']:
                        part_to_vcs.setdefault(part, set()).add(vc)
                    if line:
                        part_to_line[part] = line
                        
            pbs_dict = {}
            sealant_dict = {}
            total_dict = {}
            stages_upto_sealant = ['PBS FLOAT', 'PBS TO POLISHING', 'POLISHING TO TOPCOAT', 'TOPCOAT TO WETSANDING G ROOFBLACK', 'TOPCOAT TO WETSANDING G FRESH', 'WETSANDING G TO SEALANT']
            
            if paint_summary_vc_dict:
                for svc, counts in paint_summary_vc_dict.items():
                    p = vc_to_part.get(svc)
                    if p and str(p).strip() not in ['None', 'nan', '0']:
                        p_clean = str(p).strip()
                        total_dict[p_clean] = total_dict.get(p_clean, 0) + counts['TOTAL FLOAT']
                        pbs_dict[p_clean] = pbs_dict.get(p_clean, 0) + counts['PBS FLOAT']
                        sealant_dict[p_clean] = sealant_dict.get(p_clean, 0) + counts['TOTAL UPTO SEALANT']
            elif float_df is not None and not float_df.empty:
                vc_col = 'VEHICLE CODE' if 'VEHICLE CODE' in float_df.columns else ('VC' if 'VC' in float_df.columns else None)
                if vc_col:
                    for idx, row in float_df.iterrows():
                        vc = str(row.get(vc_col, '')).strip()[:9]
                        p = vc_to_part.get(vc)
                        if p and str(p).strip() not in ['None', 'nan', '0']:
                            p_clean = str(p).strip()
                            total_dict[p_clean] = total_dict.get(p_clean, 0) + 1
                            stg = row.get('Stage', '')
                            if stg == 'PBS FLOAT':
                                pbs_dict[p_clean] = pbs_dict.get(p_clean, 0) + 1
                            if stg in stages_upto_sealant:
                                sealant_dict[p_clean] = sealant_dict.get(p_clean, 0) + 1
                            
            today_vin_dict = {}
            for vgl_df in [tcf1_drops, tcf2_drops]:
                if vgl_df is not None and not vgl_df.empty:
                    vc_col = 'VEHICLE CODE' if 'VEHICLE CODE' in vgl_df.columns else ('VC' if 'VC' in vgl_df.columns else None)
                    if vc_col:
                        mapped_parts = vgl_df[vc_col].astype(str).str.strip().str[:9].map(vc_to_part)
                        for idx, part in mapped_parts.items():
                            if pd.notna(part) and str(part).strip() not in ['None', 'nan', '0', '']:
                                p_str = str(part).strip()
                                cnt = int(vgl_df.loc[idx, 'VIN_Count']) if 'VIN_Count' in vgl_df.columns and pd.notna(vgl_df.loc[idx, 'VIN_Count']) else 1
                                today_vin_dict[p_str] = today_vin_dict.get(p_str, 0) + cnt

            table_rows = []
            header_part_name = 'Cockpit Part Number' if 'Cockpit' in part_col_name else 'Wiring Part Number'
            
            stk_1 = stock_tcf1 if stock_tcf1 is not None else {}
            stk_2 = stock_tcf2 if stock_tcf2 is not None else {}

            for part, mdls in part_to_models.items():
                line = part_to_line.get(part, 'TCF1')
                stock_dict = stk_1 if line == 'TCF1' else stk_2
                if stock_dict and part not in stock_dict:
                    continue

                stock = stock_dict.get(part, 0)
                today_vin = today_vin_dict.get(part, 0)
                pbs = pbs_dict.get(part, 0)
                sealant = sealant_dict.get(part, 0)
                total = total_dict.get(part, 0)
                
                sh_pbs = stock - today_vin - pbs
                sh_sealant = stock - today_vin - sealant
                sh_total = stock - today_vin - total
                
                # Filter based on only_shortage parameter
                if not only_shortage or (sh_pbs < 0 or sh_sealant < 0 or sh_total < 0):
                    table_rows.append({
                        header_part_name: part,
                        'VC Number': ', '.join(sorted(part_to_vcs.get(part, []))),
                        'Model': ', '.join(sorted(mdls)),
                        'LINE': line,
                        'Clearance After 6:30AM': stock,
                        'Today VIN': today_vin,
                        'Paint TOTAL FLOAT': total,
                        'PBS FLOAT': pbs,
                        'Cabs Float UPTO SEALANT': sealant,
                        'Shortage PBS FLOAT': sh_pbs,
                        'Shortage Upto Sealant': sh_sealant,
                        'Shortage TOTAL FLOAT': sh_total
                    })
                
            return pd.DataFrame(table_rows)

        df_cpt_shortage = build_formatted_shortage_table('Cockpit', tcf1_cockpit_start, tcf2_cockpit_start, bom_df, temp_float_df, paint_summary_vc_dict, tcf1_drops, tcf2_drops, only_shortage=True)
        df_wir_shortage = build_formatted_shortage_table('Front Wiring', tcf1_wiring_start, tcf2_wiring_start, bom_df, temp_float_df, paint_summary_vc_dict, tcf1_drops, tcf2_drops, only_shortage=True)
        df_cpt_all = build_formatted_shortage_table('Cockpit', tcf1_cockpit_start, tcf2_cockpit_start, bom_df, temp_float_df, paint_summary_vc_dict, tcf1_drops, tcf2_drops, only_shortage=False)
        df_wir_all = build_formatted_shortage_table('Front Wiring', tcf1_wiring_start, tcf2_wiring_start, bom_df, temp_float_df, paint_summary_vc_dict, tcf1_drops, tcf2_drops, only_shortage=False)

        # Evaluate Today VIN excess alerts across Engine, Nova Aggregates, Cockpit, and Wiring
        excess_alerts = []
        
        # 1. Engine
        for r_eng in table2_rows:
            if r_eng.get('Type') == 'row':
                part_no = r_eng.get('Engine Part No', '')
                model_name = r_eng.get('Model', '')
                cl_val = r_eng.get('Clearance After 6:30AM')
                vin_val = r_eng.get('Today VIN', 0)
                if isinstance(cl_val, (int, float)) and vin_val > cl_val:
                    excess_alerts.append({
                        'Category': 'Engine',
                        'Model / Part': f"{model_name} ({part_no})" if part_no else model_name,
                        'Clearance 6:30 AM': cl_val,
                        'Today VIN': vin_val,
                        'Excess Qty': vin_val - cl_val
                    })
                    
        # 2. Nova Aggregates
        nova_df_check = st.session_state.get('nova_materials_df')
        if nova_df_check is not None and not nova_df_check.empty:
            vin_nova = today_vin_dict.get("546816111212", 0)
            for idx, r_n in nova_df_check.iterrows():
                m_name = r_n.get('Material', 'Aggregate')
                c_q = int(r_n.get('Clearance Qty', 0))
                if vin_nova > c_q:
                    excess_alerts.append({
                        'Category': 'Nova Aggregate',
                        'Model / Part': f"Punch EV - {m_name}",
                        'Clearance 6:30 AM': c_q,
                        'Today VIN': vin_nova,
                        'Excess Qty': vin_nova - c_q
                    })
                    
        # 3. Cockpit
        if df_cpt_all is not None and not df_cpt_all.empty:
            for idx, r_c in df_cpt_all.iterrows():
                p_hdr = 'Cockpit Part Number'
                c_no = r_c.get(p_hdr, '')
                m_descr = r_c.get('Model', '')
                cl_c = r_c.get('Clearance After 6:30AM', 0)
                vin_c = r_c.get('Today VIN', 0)
                if isinstance(cl_c, (int, float)) and vin_c > cl_c:
                    excess_alerts.append({
                        'Category': 'Cockpit',
                        'Model / Part': f"{c_no} ({m_descr})",
                        'Clearance 6:30 AM': cl_c,
                        'Today VIN': vin_c,
                        'Excess Qty': vin_c - cl_c
                    })
                    
        # 4. Wiring
        if df_wir_all is not None and not df_wir_all.empty:
            for idx, r_w in df_wir_all.iterrows():
                p_hdr = 'Wiring Part Number'
                w_no = r_w.get(p_hdr, '')
                m_descr = r_w.get('Model', '')
                cl_w = r_w.get('Clearance After 6:30AM', 0)
                vin_w = r_w.get('Today VIN', 0)
                if isinstance(cl_w, (int, float)) and vin_w > cl_w:
                    excess_alerts.append({
                        'Category': 'Wiring',
                        'Model / Part': f"{w_no} ({m_descr})",
                        'Clearance 6:30 AM': cl_w,
                        'Today VIN': vin_w,
                        'Excess Qty': vin_w - cl_w
                    })




        def render_html_formatted_shortage(df, part_header_name, is_dark):
            if df.empty:
                return "<p style='color: #6B7280; font-style: italic;'>No data available.</p>"
                
            th_bg_orange = "#382315" if is_dark else "#FCE4D6"
            th_text_orange = "#FAFAFA" if is_dark else "#73330D"
            
            th_bg_blue = "#1A2B4C" if is_dark else "#BDD7EE"
            th_text_blue = "#FAFAFA" if is_dark else "#1A2B4C"

            th_bg_blue2 = "#1E3A5F" if is_dark else "#9BC2E6"
            
            td_border = "#30363D" if is_dark else "#E5E7EB"
            text_color = "#FAFAFA" if is_dark else "#111827"
            
            alert_bg = "#5c1d1d" if is_dark else "#FFD1D1"
            alert_text = "#FAFAFA" if is_dark else "#5C1D1B"
            
            html = f"""
            <div style="overflow-x: auto; border: 1px solid {td_border}; border-radius: 12px; margin-bottom: 2rem; background-color: {'#161B22' if is_dark else '#FFFFFF'}; box-shadow: 0 4px 12px rgba(0,0,0,0.03);">
            <table style="width: 100%; border-collapse: collapse; font-family: 'Inter', sans-serif; font-size: 12px; color: {text_color};">
                <thead>
                    <tr style="border-bottom: 2px solid {td_border};">
                        <th style="padding: 10px 8px; border: 1px solid {td_border}; text-align: center; background-color: {th_bg_orange}; color: {th_text_orange}; font-weight: bold;">{part_header_name}</th>
                        <th style="padding: 10px 8px; border: 1px solid {td_border}; text-align: left; background-color: {th_bg_orange}; color: {th_text_orange}; font-weight: bold; width: 220px;">Model</th>
                        <th style="padding: 10px 8px; border: 1px solid {td_border}; text-align: center; background-color: {th_bg_orange}; color: {th_text_orange}; font-weight: bold;">LINE</th>
                        <th style="padding: 10px 8px; border: 1px solid {td_border}; text-align: center; background-color: {th_bg_blue}; color: {th_text_blue}; font-weight: bold;">Clearance After 6:30AM</th>
                        <th style="padding: 10px 8px; border: 1px solid {td_border}; text-align: center; background-color: {th_bg_blue}; color: {th_text_blue}; font-weight: bold;">Today VIN</th>
                        <th style="padding: 10px 8px; border: 1px solid {td_border}; text-align: center; background-color: {th_bg_blue}; color: {th_text_blue}; font-weight: bold;">Paint TOTAL FLOAT</th>
                        <th style="padding: 10px 8px; border: 1px solid {td_border}; text-align: center; background-color: {th_bg_blue}; color: {th_text_blue}; font-weight: bold;">PBS FLOAT</th>
                        <th style="padding: 10px 8px; border: 1px solid {td_border}; text-align: center; background-color: {th_bg_blue}; color: {th_text_blue}; font-weight: bold;">Cabs Float UPTO SEALANT</th>
                        <th style="padding: 10px 8px; border: 1px solid {td_border}; text-align: center; background-color: {th_bg_blue2}; color: {th_text_blue}; font-weight: bold;">Shortage PBS FLOAT</th>
                        <th style="padding: 10px 8px; border: 1px solid {td_border}; text-align: center; background-color: {th_bg_blue2}; color: {th_text_blue}; font-weight: bold;">Shortage Upto Sealant</th>
                        <th style="padding: 10px 8px; border: 1px solid {td_border}; text-align: center; background-color: {th_bg_blue2}; color: {th_text_blue}; font-weight: bold;">Shortage TOTAL FLOAT</th>
                    </tr>
                </thead>
                <tbody>
            """
            for idx, r in df.iterrows():
                html += f'<tr style="border-bottom: 1px solid {td_border};">'
                html += f'<td style="padding: 8px; border: 1px solid {td_border}; text-align: center; font-weight: 600;">{r[part_header_name]}</td>'
                html += f'<td style="padding: 8px; border: 1px solid {td_border}; text-align: left;">{r["Model"]}</td>'
                html += f'<td style="padding: 8px; border: 1px solid {td_border}; text-align: center;">{r["LINE"]}</td>'
                html += f'<td style="padding: 8px; border: 1px solid {td_border}; text-align: center;">{r["Clearance After 6:30AM"]}</td>'
                html += f'<td style="padding: 8px; border: 1px solid {td_border}; text-align: center;">{r["Today VIN"]}</td>'
                html += f'<td style="padding: 8px; border: 1px solid {td_border}; text-align: center;">{r["Paint TOTAL FLOAT"]}</td>'
                html += f'<td style="padding: 8px; border: 1px solid {td_border}; text-align: center;">{r["PBS FLOAT"]}</td>'
                html += f'<td style="padding: 8px; border: 1px solid {td_border}; text-align: center;">{r["Cabs Float UPTO SEALANT"]}</td>'
                
                for col_sh in ["Shortage PBS FLOAT", "Shortage Upto Sealant", "Shortage TOTAL FLOAT"]:
                    val_sh = r[col_sh]
                    if isinstance(val_sh, (int, float)) and val_sh < 0:
                        html += f'<td style="padding: 8px; border: 1px solid {td_border}; text-align: center; background-color: {alert_bg}; color: {alert_text}; font-weight: bold;">{val_sh}</td>'
                    else:
                        html += f'<td style="padding: 8px; border: 1px solid {td_border}; text-align: center;">{val_sh}</td>'
                        
                html += '</tr>'
            html += "</tbody></table></div>"
            return html


        
        # Excel generator with beautiful color schemes matching attached copy
        import io
        from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
        import openpyxl.utils
        
        excel_buffer = io.BytesIO()
        with pd.ExcelWriter(excel_buffer, engine='openpyxl') as writer:
            # Sheet 1: Summary Report
            summary_df.to_excel(writer, index=False, sheet_name='Summary Report')
            workbook = writer.book
            worksheet = writer.sheets['Summary Report']
            
            font_header = Font(name='Calibri', size=11, bold=True, color='000000')
            fill_header = PatternFill(start_color='FCE4D6', end_color='FCE4D6', fill_type='solid') # Peach
            
            font_subtotal = Font(name='Calibri', size=11, bold=True, color='000000')
            fill_subtotal = PatternFill(start_color='F2DCDB', end_color='F2DCDB', fill_type='solid') # Pink/Lavender
            
            font_grand_total = Font(name='Calibri', size=11, bold=True, color='000000')
            fill_grand_total = PatternFill(start_color='FFFF00', end_color='FFFF00', fill_type='solid') # Yellow
            
            font_normal = Font(name='Calibri', size=11, color='000000')
            
            thin_border = Border(
                left=Side(style='thin', color='BFBFBF'),
                right=Side(style='thin', color='BFBFBF'),
                top=Side(style='thin', color='BFBFBF'),
                bottom=Side(style='thin', color='BFBFBF')
            )
            
            for col_idx in range(1, len(summary_df.columns) + 1):
                cell = worksheet.cell(row=1, column=col_idx)
                cell.font = font_header
                cell.fill = fill_header
                cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
                cell.border = thin_border
                
            for row_idx in range(2, len(summary_df) + 2):
                model_val = str(worksheet.cell(row=row_idx, column=2).value).strip()
                is_subtotal = 'TOTAL' in model_val and 'GRAND' not in model_val
                is_grand = 'GRAND TOTAL' in model_val
                
                for col_idx in range(1, len(summary_df.columns) + 1):
                    cell = worksheet.cell(row=row_idx, column=col_idx)
                    cell.border = thin_border
                    cell.alignment = Alignment(horizontal='center', vertical='center')
                    
                    if is_subtotal:
                        cell.font = font_subtotal
                        cell.fill = fill_subtotal
                    elif is_grand:
                        cell.font = font_grand_total
                        cell.fill = fill_grand_total
                    else:
                        cell.font = font_normal
                        
            for col in worksheet.columns:
                max_len = max(len(str(cell.value or '')) for cell in col)
                col_letter = openpyxl.utils.get_column_letter(col[0].column)
                worksheet.column_dimensions[col_letter].width = max(max_len + 3, 12)
                
            worksheet.row_dimensions[1].height = 28
            for row_idx in range(2, len(summary_df) + 2):
                worksheet.row_dimensions[row_idx].height = 20
                
            # Sheet 2: Engine & Battery Requirement Summary
            worksheet2 = workbook.create_sheet('Engine & Battery Requirement')
            worksheet2.row_dimensions[1].height = 25
            worksheet2.row_dimensions[2].height = 25
            
            worksheet2.merge_cells('A1:A2')
            worksheet2.merge_cells('B1:B2')
            worksheet2.merge_cells('C1:C2')
            worksheet2.merge_cells('D1:D2')
            worksheet2.merge_cells('E1:E2')
            worksheet2.merge_cells('F1:F2')
            worksheet2.merge_cells('G1:I1')
            worksheet2.merge_cells('J1:L1')
            
            worksheet2['A1'] = "Engine / Battery Part No"
            worksheet2['B1'] = "Model"
            worksheet2['C1'] = "TA Code"
            worksheet2['D1'] = "Clearance After 6:30AM"
            worksheet2['E1'] = "Today VIN"
            worksheet2['F1'] = "Bal"
            worksheet2['G1'] = "Paint Float"
            worksheet2['J1'] = "Engine & Battery requirement"
            
            worksheet2['G2'] = "PBS FLOAT"
            worksheet2['H2'] = "Float UPTO SEALANT"
            worksheet2['I2'] = "TOTAL FLOAT"
            worksheet2['J2'] = "With respect to PBS FLOAT"
            worksheet2['K2'] = "With respect to Sealant FLOAT"
            worksheet2['L2'] = "With respect to Total FLOAT"
            
            for r in [1, 2]:
                for c in range(1, 13):
                    cell = worksheet2.cell(row=r, column=c)
                    cell.font = font_header
                    if r == 1 and c == 6:
                        cell.fill = PatternFill(start_color='F2DCDB', end_color='F2DCDB', fill_type='solid') # Purple/Pink
                    else:
                        cell.fill = fill_header
                    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
                    cell.border = thin_border
                    
            for r_idx, row_d in enumerate(table2_rows, start=3):
                worksheet2.row_dimensions[r_idx].height = 20
                row_type = row_d['Type']
                
                fill_row = None
                font_row = font_normal
                
                if row_type == 'subtotal':
                    fill_row = PatternFill(start_color='00B0F0', end_color='00B0F0', fill_type='solid') # Blue
                    font_row = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
                elif row_type == 'total':
                    fill_row = PatternFill(start_color='FFFF00', end_color='FFFF00', fill_type='solid') # Yellow
                    font_row = Font(name='Calibri', size=11, bold=True, color='000000')
                    
                columns_list = [
                    'Engine Part No', 'Model', 'TA Code', 'Clearance After 6:30AM',
                    'Today VIN', 'Bal', 'PBS FLOAT', 'Float UPTO SEALANT', 'TOTAL FLOAT',
                    'With respect to PBS FLOAT', 'With respect to Sealant FLOAT', 'With respect to Total FLOAT'
                ]
                
                for c_idx, col_key in enumerate(columns_list, start=1):
                    cell = worksheet2.cell(row=r_idx, column=c_idx)
                    val = row_d[col_key]
                    cell.value = val
                    cell.border = thin_border
                    cell.font = font_row
                    cell.alignment = Alignment(horizontal='center', vertical='center')
                    
                    if row_type == 'row':
                        if col_key == 'Clearance After 6:30AM':
                            cell.fill = PatternFill(start_color='D8F3E5', end_color='D8F3E5', fill_type='solid')
                            cell.font = Font(name='Calibri', size=11, bold=True, color='1B4D32')
                        elif col_key == 'Bal':
                            cell.fill = PatternFill(start_color='F2DCDB', end_color='F2DCDB', fill_type='solid')
                            cell.font = Font(name='Calibri', size=11, bold=True, color='5C1D1B')
                        elif col_key in ['With respect to PBS FLOAT', 'With respect to Sealant FLOAT', 'With respect to Total FLOAT']:
                            if isinstance(val, (int, float)) and val < 0:
                                cell.fill = PatternFill(start_color='FFD1D1', end_color='FFD1D1', fill_type='solid')
                                cell.font = Font(name='Calibri', size=11, bold=True, color='5C1D1B')
                    elif fill_row:
                        cell.fill = fill_row
                        
            for col in worksheet2.columns:
                max_len = max(len(str(cell.value or '')) for cell in col)
                col_letter = openpyxl.utils.get_column_letter(col[0].column)
                worksheet2.column_dimensions[col_letter].width = max(max_len + 3, 12)

            # Function to format openpyxl sheet for Cockpit / Wiring Shortage or All Parts
            def format_openpyxl_shortage_sheet(sheet_name, df_data, part_col_hdr, target_writer=None):
                w_target = target_writer if target_writer is not None else writer
                if df_data.empty:
                    return
                df_data.to_excel(w_target, index=False, sheet_name=sheet_name)
                ws = w_target.sheets[sheet_name]
                ws.row_dimensions[1].height = 28
                
                fill_peach = PatternFill(start_color='FCE4D6', end_color='FCE4D6', fill_type='solid')
                fill_blue = PatternFill(start_color='BDD7EE', end_color='BDD7EE', fill_type='solid')
                fill_blue2 = PatternFill(start_color='9BC2E6', end_color='9BC2E6', fill_type='solid')
                
                for c_i, col_name in enumerate(df_data.columns, start=1):
                    cell = ws.cell(row=1, column=c_i)
                    cell.font = font_header
                    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
                    cell.border = thin_border
                    if c_i <= 4:
                        cell.fill = fill_peach
                    elif c_i <= 9:
                        cell.fill = fill_blue
                    else:
                        cell.fill = fill_blue2
                        
                for r_i, r_val in enumerate(df_data.iterrows(), start=2):
                    ws.row_dimensions[r_i].height = 20
                    row_dict = r_val[1]
                    for c_i, col_name in enumerate(df_data.columns, start=1):
                        cell = ws.cell(row=r_i, column=c_i)
                        val = row_dict[col_name]
                        cell.value = val
                        cell.border = thin_border
                        cell.font = font_normal
                        cell.alignment = Alignment(horizontal='left' if col_name in ['Model', 'VC Number'] else 'center', vertical='center')
                        
                        if 'Shortage' in col_name and isinstance(val, (int, float)) and val < 0:
                            cell.fill = PatternFill(start_color='FFD1D1', end_color='FFD1D1', fill_type='solid')
                            cell.font = Font(name='Calibri', size=11, bold=True, color='5C1D1B')
                            
                for col in ws.columns:
                    max_len = max(len(str(cell.value or '')) for cell in col)
                    col_letter = openpyxl.utils.get_column_letter(col[0].column)
                    ws.column_dimensions[col_letter].width = max(max_len + 3, 14)

            format_openpyxl_shortage_sheet('Cockpit Shortage', df_cpt_shortage, 'Cockpit Part Number')
            format_openpyxl_shortage_sheet('Wiring Shortage', df_wir_shortage, 'Wiring Part Number')
            
            # Sheet: Hourly Production (if available)
            if hourly_df is not None and not hourly_df.empty:
                try:
                    df_export_hourly.to_excel(writer, index=False, sheet_name='Hourly Production')
                    ws_h_comb = writer.sheets['Hourly Production']
                    ws_h_comb.row_dimensions[1].height = 28
                    
                    fill_hdr_h = PatternFill(start_color='FCE4D6', end_color='FCE4D6', fill_type='solid')
                    fill_sub_vh = PatternFill(start_color='BDD7EE', end_color='BDD7EE', fill_type='solid')
                    fill_sub_dh = PatternFill(start_color='C6EFCE', end_color='C6EFCE', fill_type='solid')
                    
                    for c_i in range(1, len(df_export_hourly.columns) + 1):
                        c_cell = ws_h_comb.cell(row=1, column=c_i)
                        c_cell.font = font_header
                        c_cell.fill = fill_hdr_h
                        c_cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
                        c_cell.border = thin_border
                        
                    for r_i in range(2, len(df_export_hourly) + 2):
                        ws_h_comb.row_dimensions[r_i].height = 22
                        pt_val = str(ws_h_comb.cell(row=r_i, column=1).value).strip().upper()
                        is_tot_v = 'TOTAL VIN' in pt_val
                        is_tot_d = 'TOTAL TCF DROP' in pt_val or 'TOTAL DROP' in pt_val
                        
                        for c_i in range(1, len(df_export_hourly.columns) + 1):
                            cell_obj = ws_h_comb.cell(row=r_i, column=c_i)
                            cell_obj.border = thin_border
                            if c_i == 1:
                                cell_obj.alignment = Alignment(horizontal='left', vertical='center')
                            else:
                                cell_obj.alignment = Alignment(horizontal='center', vertical='center')
                                
                            if is_tot_v:
                                cell_obj.font = font_subtotal
                                cell_obj.fill = fill_sub_vh
                            elif is_tot_d:
                                cell_obj.font = font_subtotal
                                cell_obj.fill = fill_sub_dh
                            else:
                                cell_obj.font = font_normal
                                
                    for col in ws_h_comb.columns:
                        m_len = max(len(str(cell.value or '')) for cell in col)
                        c_let = openpyxl.utils.get_column_letter(col[0].column)
                        ws_h_comb.column_dimensions[c_let].width = max(m_len + 4, 14)
                except Exception as ex_h:
                    print(f"Error adding Hourly Production to summary export: {ex_h}")
                
        excel_data = excel_buffer.getvalue()

        # Build 2-Sheet Excel workbook for Cockpit & Wiring Report (All Parts: Cockpit, Wiring)
        all_parts_excel_buffer = io.BytesIO()
        with pd.ExcelWriter(all_parts_excel_buffer, engine='openpyxl') as writer_all:
            format_openpyxl_shortage_sheet('Cockpit', df_cpt_all, 'Cockpit Part Number', target_writer=writer_all)
            format_openpyxl_shortage_sheet('Wiring', df_wir_all, 'Wiring Part Number', target_writer=writer_all)
        all_parts_excel_data = all_parts_excel_buffer.getvalue()
        
        st.markdown("---")
        st.download_button(
            label="📥 Export Master Summary Reports to Excel",
            data=excel_data,
            file_name="paint_shop_float_and_requirements_summary.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            key="export_summary_report"
        )

        # --- FINAL REMARK: BOM COMPLETENESS ALERT & QUICK-ENTRY ---
        st.markdown("---")
        if 'missing_bom_df' in locals() and not missing_bom_df.empty:
            affected_cabs = int(missing_bom_df['Cab Count'].sum())
            st.error(
                f"🚨 **BOM not available for {len(missing_bom_df)} Short Vehicle Code(s)** "
                f"({affected_cabs} cab(s) affected). These cabs may be miscounted as Ready-to-TCF "
                "with blank/NaN parts until BOM is entered below."
            )
            with st.expander(f"⚠️ Fix Missing/Incomplete BOM — {len(missing_bom_df)} Short VC(s)", expanded=True):
                st.dataframe(missing_bom_df, use_container_width=True, hide_index=True)
                st.markdown("###### ➕ Enter BOM for a Short VC")
                with st.form("missing_bom_entry_form", clear_on_submit=True):
                    sel_vc = st.selectbox("Short Vehicle Code", options=missing_bom_df['Short VC'].tolist())
                    fb1, fb2, fb3 = st.columns(3)
                    with fb1:
                        in_engine = st.text_input("Engine / Battery Part No.")
                    with fb2:
                        in_cockpit = st.text_input("Cockpit Part No.")
                    with fb3:
                        in_wiring = st.text_input("Front Wiring Part No.")
                    submitted = st.form_submit_button("💾 Save BOM Entry")
                    if submitted:
                        if not (in_engine.strip() or in_cockpit.strip() or in_wiring.strip()):
                            st.warning("Enter at least one part number before saving.")
                        else:
                            try:
                                dl.save_single_bom_entry(sel_vc, in_wiring, in_cockpit, in_engine)
                                st.success(f"Saved BOM for {sel_vc} to the database. Refreshing report...")
                                st.session_state.run_report = True
                                st.rerun()
                            except Exception as e:
                                st.error(f"Could not save BOM entry: {e}")
        elif 'missing_bom_df' in locals() and missing_bom_df.empty and bom_df is not None and not bom_df.empty:
            st.success("✅ All BOM data checked — no error found. Every Short VC in the current float has a complete BOM match.")
    else:
        st.info("Please load Paint Float data in the Control Panel to view the summary report.")


