import streamlit as st
import pandas as pd
import numpy as np
import io
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
import openpyxl.utils
from core.data_manager import get_dashboard_data

def format_openpyxl_shortage_sheet(sheet_name, df_data, part_col_hdr, target_writer=None):
    if df_data is None or df_data.empty:
        pd.DataFrame(columns=[part_col_hdr, 'Model', 'LINE']).to_excel(target_writer, index=False, sheet_name=sheet_name)
        return
    df_data.to_excel(target_writer, index=False, sheet_name=sheet_name)
    worksheet = target_writer.sheets[sheet_name]
    
    font_header = Font(name='Calibri', size=11, bold=True, color='000000')
    font_normal = Font(name='Calibri', size=11, color='000000')
    font_alert = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
    
    fill_header = PatternFill(start_color='FCE4D6', end_color='FCE4D6', fill_type='solid')
    fill_header_blue = PatternFill(start_color='BDD7EE', end_color='BDD7EE', fill_type='solid')
    fill_header_blue2 = PatternFill(start_color='9BC2E6', end_color='9BC2E6', fill_type='solid')
    fill_alert = PatternFill(start_color='FF0000', end_color='FF0000', fill_type='solid')
    
    thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
    
    for idx, col_name in enumerate(df_data.columns, 1):
        cell = worksheet.cell(row=1, column=idx)
        cell.font = font_header
        cell.border = thin_border
        if 'Shortage' in col_name:
            cell.fill = fill_header_blue2
        elif col_name in ['Clearance After 6:30AM', 'Today VIN', 'Paint TOTAL FLOAT', 'PBS FLOAT', 'Cabs Float UPTO SEALANT']:
            cell.fill = fill_header_blue
        else:
            cell.fill = fill_header
            
    for row_idx in range(2, len(df_data) + 2):
        for col_idx, col_name in enumerate(df_data.columns, 1):
            cell = worksheet.cell(row=row_idx, column=col_idx)
            cell.font = font_normal
            cell.border = thin_border
            if 'Shortage' in col_name and isinstance(cell.value, (int, float)) and cell.value < 0:
                cell.fill = fill_alert
                cell.font = font_alert
                
    for col in worksheet.columns:
        max_length = 0
        col_letter = col[0].column_letter
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        worksheet.column_dimensions[col_letter].width = min(max_length + 2, 40)

def render_html_formatted_shortage(df, part_header_name, is_dark):
    if df.empty:
        return "<p style='color: #6B7280; font-style: italic;'>No data available.</p>"
        
    th_bg_orange = "#382315" if is_dark else "#FCE4D6"
    th_text_orange = "#FAFAFA" if is_dark else "#73330D"
    th_bg_blue = "#1A2B4C" if is_dark else "#BDD7EE"
    th_text_blue = "#FAFAFA" if is_dark else "#1A2B4C"
    th_bg_blue2 = "#1E3A5F" if is_dark else "#9BC2E6"
    td_border = "#30363D" if is_dark else "#E5E7EB"
    text_color = "#FAFAFA" if is_dark else "#111827"
    alert_bg = "#5c1d1d" if is_dark else "#FFD1D1"
    alert_text = "#FAFAFA" if is_dark else "#5C1D1B"
    bg_table = '#161B22' if is_dark else '#FFFFFF'
    
    html = f"""
    <div style="overflow-x: auto; border: 1px solid {td_border}; border-radius: 12px; margin-bottom: 2rem; background-color: {bg_table}; box-shadow: 0 4px 12px rgba(0,0,0,0.03);">
    <table style="width: 100%; border-collapse: collapse; font-family: 'Inter', sans-serif; font-size: 12px; color: {text_color};">
        <thead>
            <tr style="border-bottom: 2px solid {td_border};">
                <th style="padding: 10px 8px; border: 1px solid {td_border}; text-align: center; background-color: {th_bg_orange}; color: {th_text_orange}; font-weight: bold;">{part_header_name}</th>
                <th style="padding: 10px 8px; border: 1px solid {td_border}; text-align: left; background-color: {th_bg_orange}; color: {th_text_orange}; font-weight: bold; width: 220px;">Model</th>
                <th style="padding: 10px 8px; border: 1px solid {td_border}; text-align: center; background-color: {th_bg_orange}; color: {th_text_orange}; font-weight: bold;">LINE</th>
                <th style="padding: 10px 8px; border: 1px solid {td_border}; text-align: center; background-color: {th_bg_blue}; color: {th_text_blue}; font-weight: bold;">Clearance After 6:30AM</th>
                <th style="padding: 10px 8px; border: 1px solid {td_border}; text-align: center; background-color: {th_bg_blue}; color: {th_text_blue}; font-weight: bold;">Today VIN</th>
                <th style="padding: 10px 8px; border: 1px solid {td_border}; text-align: center; background-color: {th_bg_blue}; color: {th_text_blue}; font-weight: bold;">Paint TOTAL FLOAT</th>
                <th style="padding: 10px 8px; border: 1px solid {td_border}; text-align: center; background-color: {th_bg_blue}; color: {th_text_blue}; font-weight: bold;">PBS FLOAT</th>
                <th style="padding: 10px 8px; border: 1px solid {td_border}; text-align: center; background-color: {th_bg_blue}; color: {th_text_blue}; font-weight: bold;">Cabs Float UPTO SEALANT</th>
                <th style="padding: 10px 8px; border: 1px solid {td_border}; text-align: center; background-color: {th_bg_blue2}; color: {th_text_blue}; font-weight: bold;">Shortage PBS FLOAT</th>
                <th style="padding: 10px 8px; border: 1px solid {td_border}; text-align: center; background-color: {th_bg_blue2}; color: {th_text_blue}; font-weight: bold;">Shortage Upto Sealant</th>
                <th style="padding: 10px 8px; border: 1px solid {td_border}; text-align: center; background-color: {th_bg_blue2}; color: {th_text_blue}; font-weight: bold;">Shortage TOTAL FLOAT</th>
            </tr>
        </thead>
        <tbody>
    """
    for idx, r in df.iterrows():
        html += f'<tr style="border-bottom: 1px solid {td_border};">'
        html += f'<td style="padding: 8px; border: 1px solid {td_border}; text-align: center; font-weight: 600;">{r[part_header_name]}</td>'
        html += f'<td style="padding: 8px; border: 1px solid {td_border}; text-align: left;">{r["Model"]}</td>'
        html += f'<td style="padding: 8px; border: 1px solid {td_border}; text-align: center;">{r["LINE"]}</td>'
        html += f'<td style="padding: 8px; border: 1px solid {td_border}; text-align: center;">{r["Clearance After 6:30AM"]}</td>'
        html += f'<td style="padding: 8px; border: 1px solid {td_border}; text-align: center;">{r["Today VIN"]}</td>'
        html += f'<td style="padding: 8px; border: 1px solid {td_border}; text-align: center;">{r["Paint TOTAL FLOAT"]}</td>'
        html += f'<td style="padding: 8px; border: 1px solid {td_border}; text-align: center;">{r["PBS FLOAT"]}</td>'
        html += f'<td style="padding: 8px; border: 1px solid {td_border}; text-align: center;">{r["Cabs Float UPTO SEALANT"]}</td>'
        
        for col_sh in ["Shortage PBS FLOAT", "Shortage Upto Sealant", "Shortage TOTAL FLOAT"]:
            val_sh = r[col_sh]
            if isinstance(val_sh, (int, float)) and val_sh < 0:
                html += f'<td style="padding: 8px; border: 1px solid {td_border}; text-align: center; background-color: {alert_bg}; color: {alert_text}; font-weight: bold;">{val_sh}</td>'
            else:
                html += f'<td style="padding: 8px; border: 1px solid {td_border}; text-align: center;">{val_sh}</td>'
                
        html += '</tr>'
    html += "</tbody></table></div>"
    return html

def build_formatted_shortage_table(part_col_name, stock_tcf1, stock_tcf2, bom_df, float_df, paint_summary_vc_dict, tcf1_drops, tcf2_drops, only_shortage=True):
    if bom_df is None or bom_df.empty:
        return pd.DataFrame()
        
    local_engine_to_model = {}
    local_engine_to_line = {}
    if 'engine_df' in st.session_state and not st.session_state.engine_df.empty:
        local_engine_to_model = dict(zip(st.session_state.engine_df['Engine Part No'].astype(str).str.strip(), st.session_state.engine_df['Model']))
        local_engine_to_line = dict(zip(st.session_state.engine_df['Engine Part No'].astype(str).str.strip(), st.session_state.engine_df['TCF Line']))
    else:
        engine_default_data = [
            {"TCF Line": "TCF1", "Engine Part No": "54850000PTP001", "Model": "Punch MT SA"},
            {"TCF Line": "TCF1", "Engine Part No": "54850000PTP002", "Model": "Punch AMT SA"},
            {"TCF Line": "TCF1", "Engine Part No": "54970000PTP002", "Model": "Punch TC MCE"},
            {"TCF Line": "TCF1", "Engine Part No": "54970000PTP003", "Model": "Punch MCE MT"},
            {"TCF Line": "TCF1", "Engine Part No": "54970000PTP004", "Model": "Punch MCE AMT"},
            {"TCF Line": "TCF1", "Engine Part No": "54970000PTP005", "Model": "Punch MCE CNG MT"},
            {"TCF Line": "TCF1", "Engine Part No": "54970000PTP031", "Model": "Punch MCE CNG AMT"},
            {"TCF Line": "TCF2", "Engine Part No": "572900000118", "Model": "Harrier / Safari Diesel AT"},
            {"TCF Line": "TCF2", "Engine Part No": "572900000120", "Model": "Harrier / Safari Diesel MT"},
            {"TCF Line": "TCF2", "Engine Part No": "54780000PTP001", "Model": "Harrier / Safari Petrol TGDI MT"},
            {"TCF Line": "TCF2", "Engine Part No": "54780000PTP002", "Model": "Harrier / Safari Petrol TGDI AT"}
        ]
        local_engine_to_model = {item['Engine Part No']: item['Model'] for item in engine_default_data}
        local_engine_to_line = {item['Engine Part No']: item['TCF Line'] for item in engine_default_data}

    local_engine_to_model['546816111212'] = 'Punch EV (Nova)'
    local_engine_to_line['546816111212'] = 'TCF1'
    local_engine_to_model['547380400103'] = 'Harrier EV'
    local_engine_to_line['547380400103'] = 'TCF2'

    vc_to_part = dict(zip(bom_df['Short Vehicle Code'].astype(str).str.strip(), bom_df[part_col_name].astype(str).str.strip()))
    
    part_to_models = {}
    part_to_vcs = {}
    part_to_line = {}
    
    for idx, row in bom_df.iterrows():
        eng = str(row.get('Engine', '')).strip()
        part = str(row.get(part_col_name, '')).strip()
        vc = str(row.get('Short Vehicle Code', '')).strip()
        mdl = local_engine_to_model.get(eng, '')
        line = local_engine_to_line.get(eng, '')
        if part and part not in ['None', 'nan', '0']:
            if mdl:
                part_to_models.setdefault(part, set()).add(mdl)
            if vc and vc not in ['None', 'nan', '0']:
                part_to_vcs.setdefault(part, set()).add(vc)
            if line:
                part_to_line[part] = line
                
    pbs_dict = {}
    sealant_dict = {}
    total_dict = {}
    stages_upto_sealant = ['PBS FLOAT', 'PBS TO POLISHING', 'POLISHING TO TOPCOAT', 'TOPCOAT TO WETSANDING G ROOFBLACK', 'TOPCOAT TO WETSANDING G FRESH', 'WETSANDING G TO SEALANT']
    
    if paint_summary_vc_dict:
        for svc, counts in paint_summary_vc_dict.items():
            p = vc_to_part.get(svc)
            if p and str(p).strip() not in ['None', 'nan', '0']:
                p_clean = str(p).strip()
                total_dict[p_clean] = total_dict.get(p_clean, 0) + counts.get('TOTAL FLOAT', 0)
                pbs_dict[p_clean] = pbs_dict.get(p_clean, 0) + counts.get('PBS FLOAT', 0)
                sealant_dict[p_clean] = sealant_dict.get(p_clean, 0) + counts.get('TOTAL UPTO SEALANT', 0)
    elif float_df is not None and not float_df.empty:
        vc_col = 'VEHICLE CODE' if 'VEHICLE CODE' in float_df.columns else ('VC' if 'VC' in float_df.columns else None)
        if vc_col:
            for idx, row in float_df.iterrows():
                vc = str(row.get(vc_col, '')).strip()[:9]
                p = vc_to_part.get(vc)
                if p and str(p).strip() not in ['None', 'nan', '0']:
                    p_clean = str(p).strip()
                    total_dict[p_clean] = total_dict.get(p_clean, 0) + 1
                    stg = row.get('Stage', '')
                    if stg == 'PBS FLOAT':
                        pbs_dict[p_clean] = pbs_dict.get(p_clean, 0) + 1
                    if stg in stages_upto_sealant:
                        sealant_dict[p_clean] = sealant_dict.get(p_clean, 0) + 1
                    
    today_vin_dict = {}
    for vgl_df in [tcf1_drops, tcf2_drops]:
        if vgl_df is not None and not vgl_df.empty:
            vc_col = 'VEHICLE CODE' if 'VEHICLE CODE' in vgl_df.columns else ('VC' if 'VC' in vgl_df.columns else None)
            if vc_col:
                mapped_parts = vgl_df[vc_col].astype(str).str.strip().str[:9].map(vc_to_part)
                for idx, part in mapped_parts.items():
                    if pd.notna(part) and str(part).strip() not in ['None', 'nan', '0', '']:
                        p_str = str(part).strip()
                        cnt = int(vgl_df.loc[idx, 'VIN_Count']) if 'VIN_Count' in vgl_df.columns and pd.notna(vgl_df.loc[idx, 'VIN_Count']) else 1
                        today_vin_dict[p_str] = today_vin_dict.get(p_str, 0) + cnt

    table_rows = []
    header_part_name = 'Cockpit Part Number' if 'Cockpit' in part_col_name else 'Wiring Part Number'
    
    stk_1 = stock_tcf1 if stock_tcf1 is not None else {}
    stk_2 = stock_tcf2 if stock_tcf2 is not None else {}

    for part, mdls in part_to_models.items():
        line = part_to_line.get(part, 'TCF1')
        stock_dict = stk_1 if line == 'TCF1' else stk_2
        if stock_dict and part not in stock_dict:
            continue

        stock = stock_dict.get(part, 0)
        today_vin = today_vin_dict.get(part, 0)
        pbs = pbs_dict.get(part, 0)
        sealant = sealant_dict.get(part, 0)
        total = total_dict.get(part, 0)
        
        sh_pbs = stock - today_vin - pbs
        sh_sealant = stock - today_vin - sealant
        sh_total = stock - today_vin - total
        
        if not only_shortage or (sh_pbs < 0 or sh_sealant < 0 or sh_total < 0):
            table_rows.append({
                header_part_name: part,
                'VC Number': ', '.join(sorted(part_to_vcs.get(part, []))),
                'Model': ', '.join(sorted(mdls)),
                'LINE': line,
                'Clearance After 6:30AM': stock,
                'Today VIN': today_vin,
                'Paint TOTAL FLOAT': total,
                'PBS FLOAT': pbs,
                'Cabs Float UPTO SEALANT': sealant,
                'Shortage PBS FLOAT': sh_pbs,
                'Shortage Upto Sealant': sh_sealant,
                'Shortage TOTAL FLOAT': sh_total
            })
        
    return pd.DataFrame(table_rows)

def render():
    classified_files = st.session_state.get("classified_files", {})
    st.markdown("### 🧩 Cockpit & Wiring Shortage Reports")
    st.markdown("""
        Real-time shortage monitoring for **Cockpit Assemblies** and **Front Wiring Harnesses** matching engine summary models across TCF1 and TCF2 lines.
    """)

    db_data = get_dashboard_data(classified_files)
    bom_df = db_data.bom_df
    float_df = db_data.float_df
    paint_summary_vc_dict = db_data.paint_summary
    
    tcf1_drops = db_data.vgl_df[db_data.vgl_df['SHOP'] == 'TCF1'] if not db_data.vgl_df.empty and 'SHOP' in db_data.vgl_df.columns else None
    tcf2_drops = db_data.vgl_df[db_data.vgl_df['SHOP'] == 'TCF2'] if not db_data.vgl_df.empty and 'SHOP' in db_data.vgl_df.columns else None

    # Load clearance buffers from session state (used by build_formatted_shortage_table)
    tcf1_cockpit_start = st.session_state.get("buffer_cockpit", {})
    tcf2_cockpit_start = st.session_state.get("buffer_cockpit", {}) # App.py might have different logic, but buffer_cockpit holds both
    tcf1_wiring_start = st.session_state.get("buffer_wiring", {})
    tcf2_wiring_start = st.session_state.get("buffer_wiring", {})

    df_cpt_shortage = build_formatted_shortage_table('Cockpit', tcf1_cockpit_start, tcf2_cockpit_start, bom_df, float_df, paint_summary_vc_dict, tcf1_drops, tcf2_drops, only_shortage=True)
    df_wir_shortage = build_formatted_shortage_table('Front Wiring', tcf1_wiring_start, tcf2_wiring_start, bom_df, float_df, paint_summary_vc_dict, tcf1_drops, tcf2_drops, only_shortage=True)
    df_cpt_all = build_formatted_shortage_table('Cockpit', tcf1_cockpit_start, tcf2_cockpit_start, bom_df, float_df, paint_summary_vc_dict, tcf1_drops, tcf2_drops, only_shortage=False)
    df_wir_all = build_formatted_shortage_table('Front Wiring', tcf1_wiring_start, tcf2_wiring_start, bom_df, float_df, paint_summary_vc_dict, tcf1_drops, tcf2_drops, only_shortage=False)

    cpt_sh_df = df_cpt_shortage
    wir_sh_df = df_wir_shortage
    cpt_all_df = df_cpt_all
    wir_all_df = df_wir_all

    if (cpt_all_df is None or cpt_all_df.empty) and (wir_all_df is None or wir_all_df.empty):
        st.info("ℹ️ Please load Paint Float and BOM data in the Control Panel to view Cockpit & Wiring Shortage Reports.")
    else:
        # Top KPI Summary Cards
        cpt_sh_count = len(cpt_sh_df) if cpt_sh_df is not None else 0
        wir_sh_count = len(wir_sh_df) if wir_sh_df is not None else 0
        tot_cpt_count = len(cpt_all_df) if cpt_all_df is not None else 0
        tot_wir_count = len(wir_all_df) if wir_all_df is not None else 0

        kpi_sh1, kpi_sh2, kpi_sh3, kpi_sh4 = st.columns(4)
        with kpi_sh1:
            st.metric(
                label="🚗 Cockpit Shortages",
                value=f"{cpt_sh_count} Part{'s' if cpt_sh_count != 1 else ''}",
                delta="Critical Shortage" if cpt_sh_count > 0 else "All Covered",
                delta_color="inverse" if cpt_sh_count > 0 else "normal"
            )
        with kpi_sh2:
            st.metric(
                label="⚡ Wiring Shortages",
                value=f"{wir_sh_count} Part{'s' if wir_sh_count != 1 else ''}",
                delta="Critical Shortage" if wir_sh_count > 0 else "All Covered",
                delta_color="inverse" if wir_sh_count > 0 else "normal"
            )
        with kpi_sh3:
            st.metric(
                label="📦 Monitored Cockpits",
                value=f"{tot_cpt_count} Part Numbers"
            )
        with kpi_sh4:
            st.metric(
                label="🔌 Monitored Wiring",
                value=f"{tot_wir_count} Part Numbers"
            )

        # Filters: View Mode, Line, Search
        st.markdown("<div style='height: 8px;'></div>", unsafe_allow_html=True)
        f_col1, f_col2, f_col3 = st.columns([1.5, 1.2, 1.3])
        with f_col1:
            sh_view_mode = st.radio(
                "📋 Display Filter:",
                ["🚨 Critical Shortages Only", "📋 All Parts & Current Stock"],
                horizontal=True,
                key="sh_tab_view_mode"
            )
        with f_col2:
            sh_line_filter = st.selectbox(
                "🏭 Filter by TCF Line:",
                ["All Lines (TCF1 + TCF2)", "TCF1 Line", "TCF2 Line"],
                key="sh_tab_line_filter"
            )
        with f_col3:
            sh_search = st.text_input(
                "🔍 Search Part / Model:",
                placeholder="Search part number or model...",
                key="sh_tab_search"
            )

        show_shortages_only = (sh_view_mode == "🚨 Critical Shortages Only")
        target_cpt = cpt_sh_df if show_shortages_only else cpt_all_df
        target_wir = wir_sh_df if show_shortages_only else wir_all_df

        def _apply_sh_filters(df_in, part_hdr):
            if df_in is None or df_in.empty:
                return df_in
            df_out = df_in.copy()
            if sh_line_filter == "TCF1 Line":
                df_out = df_out[df_out['LINE'] == 'TCF1']
            elif sh_line_filter == "TCF2 Line":
                df_out = df_out[df_out['LINE'] == 'TCF2']
            if sh_search.strip():
                q = sh_search.strip().lower()
                df_out = df_out[
                    df_out[part_hdr].astype(str).str.lower().str.contains(q) |
                    df_out['Model'].astype(str).str.lower().str.contains(q) |
                    df_out['VC Number'].astype(str).str.lower().str.contains(q)
                ]
            return df_out

        filtered_cpt = _apply_sh_filters(target_cpt, "Cockpit Part Number")
        filtered_wir = _apply_sh_filters(target_wir, "Wiring Part Number")

        is_dark_theme = st.session_state.get('theme', '☀️ White Theme') == '🌙 Dark Theme'

        st.markdown("---")
        st.markdown("#### 🚗 Cockpit Shortage Report")
        if filtered_cpt is not None and not filtered_cpt.empty:
            st.markdown(render_html_formatted_shortage(filtered_cpt, "Cockpit Part Number", is_dark_theme), unsafe_allow_html=True)
        else:
            if show_shortages_only:
                st.success("✅ No Cockpit Shortages detected for selected filters! All required cockpit assemblies are covered by clearance stock.")
            else:
                st.info("No cockpit records match your filters.")

        st.markdown("#### ⚡ Wiring Harness Shortage Report")
        if filtered_wir is not None and not filtered_wir.empty:
            st.markdown(render_html_formatted_shortage(filtered_wir, "Wiring Part Number", is_dark_theme), unsafe_allow_html=True)
        else:
            if show_shortages_only:
                st.success("✅ No Wiring Harness Shortages detected for selected filters! All required wiring harnesses are covered by clearance stock.")
            else:
                st.info("No wiring harness records match your filters.")

        # Dedicated Export Section in Tab 1
        st.markdown("---")
        exp_sh1, exp_sh2 = st.columns(2)
        with exp_sh1:
            all_parts_excel_buf = io.BytesIO()
            with pd.ExcelWriter(all_parts_excel_buf, engine='openpyxl') as writer_all:
                format_openpyxl_shortage_sheet('Cockpit All Parts', cpt_all_df, 'Cockpit Part Number', target_writer=writer_all)
                format_openpyxl_shortage_sheet('Wiring All Parts', wir_all_df, 'Wiring Part Number', target_writer=writer_all)
                
            st.download_button(
                label="📥 Download Cockpit & Wiring Report (All Parts - 2 Sheets)",
                data=all_parts_excel_buf.getvalue(),
                file_name="Cockpit_and_Wiring_Report_All_Parts.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                key="export_cockpit_wiring_tab_all_parts"
            )
        with exp_sh2:
            sh_excel_buf = io.BytesIO()
            with pd.ExcelWriter(sh_excel_buf, engine='openpyxl') as writer_sh_only:
                format_openpyxl_shortage_sheet('Cockpit Shortage', cpt_sh_df, 'Cockpit Part Number', target_writer=writer_sh_only)
                format_openpyxl_shortage_sheet('Wiring Shortage', wir_sh_df, 'Wiring Part Number', target_writer=writer_sh_only)
            st.download_button(
                label="📥 Download Critical Shortages Only (Excel)",
                data=sh_excel_buf.getvalue(),
                file_name="Cockpit_and_Wiring_Critical_Shortages.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                key="export_cockpit_wiring_tab_critical_only"
            )
