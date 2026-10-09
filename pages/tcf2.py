import streamlit as st
import pandas as pd
import io
import sys
import os
from pathlib import Path
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
import openpyxl.utils

project_root = Path(__file__).parent.parent
if str(project_root) not in sys.path:
    sys.path.insert(0, str(project_root))

from core.data_manager import get_dashboard_data

def render():
    classified_files = st.session_state.get('classified_files', {})
    try:
        data = get_dashboard_data(classified_files)
    except TypeError:
        data = get_dashboard_data(classified_files)

    if data.status == "error":
        st.error(f"Error loading dashboard data: {data.error_message}")
        return

    alloc_df = pd.DataFrame(data.alloc_results)
    if alloc_df.empty:
        tcf2_alloc_df = pd.DataFrame()
    else:
        tcf2_alloc_df = alloc_df[alloc_df.get('SHOP', '') == 'TCF2'].copy()

    float_df = data.float_df
    tcf2_queue = pd.DataFrame()
    pbs_on_hold = pd.DataFrame()
    if not float_df.empty and 'SHOP' in float_df.columns:
        pbs_all = float_df[float_df['PBS LIFT'].notna()].copy() if 'PBS LIFT' in float_df.columns else float_df.copy()
        if 'HOLD BY' in pbs_all.columns:
            is_hold_pbs = pbs_all['HOLD BY'].notna() & (pbs_all['HOLD BY'].astype(str).str.strip() != '') & (pbs_all['HOLD BY'].astype(str).str.upper() != 'NONE')
            pbs_on_hold = pbs_all[is_hold_pbs].copy()
            pbs_active = pbs_all[~is_hold_pbs].copy()
        else:
            pbs_on_hold = pd.DataFrame()
            pbs_active = pbs_all.copy()
        tcf2_queue = pbs_active[pbs_active['SHOP'] == 'TCF2'].copy()

    tcf2_drops = pd.DataFrame()
    if not data.vgl_df.empty and 'Platform' in data.vgl_df.columns:
        tcf2_drops = data.vgl_df[data.vgl_df['Platform'].isin(['HARRIER', 'SAFARI', 'HARRIER.EV', 'SAFARI.EV', 'Q5', 'TAYRONA', 'ETURNA'])]
    elif not data.vgl_df.empty:
        tcf2_drops = data.vgl_df

    # KPIs
    ready_count_tcf2 = len(tcf2_alloc_df[tcf2_alloc_df['STATUS'] == '✅ Ready for TCF']) if not tcf2_alloc_df.empty else 0
    blocked_count_tcf2 = len(tcf2_alloc_df[tcf2_alloc_df['STATUS'] == '🚫 Blocked']) if not tcf2_alloc_df.empty else 0
    issue_count_tcf2 = len(tcf2_alloc_df[tcf2_alloc_df['STATUS'].str.startswith('⚠️', na=False)]) if not tcf2_alloc_df.empty else 0
    total_drops_tcf2 = int(tcf2_drops['VIN_Count'].sum()) if (not tcf2_drops.empty and 'VIN_Count' in tcf2_drops.columns) else len(tcf2_drops)
    
    tcf2_ok = len(tcf2_queue)
    tcf2_hold = len(pbs_on_hold[pbs_on_hold['SHOP'] == 'TCF2']) if not pbs_on_hold.empty and 'SHOP' in pbs_on_hold.columns else 0
    tcf2_total = tcf2_ok + tcf2_hold
    
    kpi_cols_tcf2 = st.columns(4)
    kpi_cols_tcf2[0].metric("VIN Generation", f"{total_drops_tcf2} cabs", help="Cabs built in TCF2 since shift start")
    kpi_cols_tcf2[1].metric("PBS Current Stock", f"{tcf2_total} cabs (OK: {tcf2_ok} | Hold: {tcf2_hold})", help="Total cabs in TCF2 PBS buffer (Active unblocked + Quality holds)")
    kpi_cols_tcf2[2].metric("✅ Ready for TCF", f"{ready_count_tcf2} cabs", delta=f"+{ready_count_tcf2} alloc")
    kpi_cols_tcf2[3].metric("🚫 Blocked (Stock Out)", f"{blocked_count_tcf2} cabs", delta=f"-{blocked_count_tcf2} wait", delta_color="inverse")
    
    st.markdown("### TCF2 FIFO Buffer Queue Status")
    if tcf2_alloc_df.empty:
        st.info("No active cabs in TCF2 PBS queue.")
    else:
        # Filter to show only Ready for TCF and Blocked statuses
        filtered_df_tcf2 = tcf2_alloc_df[tcf2_alloc_df['STATUS'].isin(['✅ Ready for TCF', '🚫 Blocked'])].copy()
        
        # Apply any manual planner overrides from session_state
        if 'tcf2_manual_overrides' not in st.session_state:
            st.session_state.tcf2_manual_overrides = {}
        for biw_key, override in st.session_state.tcf2_manual_overrides.items():
            mask = filtered_df_tcf2['BIW NUMBER'].astype(str) == str(biw_key)
            if mask.any():
                filtered_df_tcf2.loc[mask, 'STATUS'] = override['status']
                filtered_df_tcf2.loc[mask, 'BLOCKING_REASON'] = override['reason']
        
        c_srch1_tcf2, c_srch2_tcf2 = st.columns([3, 3])
        with c_srch1_tcf2:
            search_biw_tcf2 = st.text_input("🔍 Quick Search by BIW Number:", key="tcf2_biw_search")
        with c_srch2_tcf2:
            loc_options_tcf2 = ['All Locations'] + sorted(list(filtered_df_tcf2['Cab location'].dropna().unique())) if 'Cab location' in filtered_df_tcf2.columns else ['All Locations']
            selected_loc_tcf2 = st.selectbox("📍 Cab Location Filter:", options=loc_options_tcf2, key="tcf2_loc_search")
            
        if search_biw_tcf2:
            filtered_df_tcf2 = filtered_df_tcf2[filtered_df_tcf2['BIW NUMBER'].astype(str).str.contains(search_biw_tcf2.strip())]
        if selected_loc_tcf2 != 'All Locations':
            filtered_df_tcf2 = filtered_df_tcf2[filtered_df_tcf2['Cab location'] == selected_loc_tcf2]
            
        display_cols = ['BIW NUMBER', 'Model', 'Trim', 'VEHICLE CODE', 'STATUS', 'BLOCKING_REASON', 'Cab location', 'Engine_Part', 'Cockpit_Part', 'Wiring_Part', 'Engine_Stock_After', 'Cockpit_Stock_After', 'Wiring_Stock_After']
        display_cols = [c for c in display_cols if c in filtered_df_tcf2.columns]
        
        st.caption("✏️ **Planner Edit Mode** — Click any cell in the **Status** or **Blocking Reason** column to change it. Changes are saved automatically.")
        
        status_options = ['✅ Ready for TCF', '🚫 Blocked', '⚠️ PBS Hold']
        
        edited_tcf2 = st.data_editor(
            filtered_df_tcf2[display_cols],
            use_container_width=True,
            hide_index=True,
            key="tcf2_queue_editor",
            column_config={
                "BIW NUMBER": st.column_config.TextColumn("BIW NUMBER", disabled=True),
                "Model": st.column_config.TextColumn("Model", disabled=True),
                "Trim": st.column_config.TextColumn("Trim", disabled=True),
                "VEHICLE CODE": st.column_config.TextColumn("VEHICLE CODE", disabled=True),
                "STATUS": st.column_config.SelectboxColumn(
                    "STATUS",
                    options=status_options,
                    required=True,
                    help="Select cab status: Ready, Blocked, or PBS Hold"
                ),
                "BLOCKING_REASON": st.column_config.TextColumn(
                    "BLOCKING REASON",
                    help="Enter blocking/hold reason (e.g., quality issue, part shortage, PBS hold)"
                ),
                "Cab location": st.column_config.TextColumn("Cab Location", disabled=True),
                "Engine_Part": st.column_config.TextColumn("Engine Part", disabled=True),
                "Cockpit_Part": st.column_config.TextColumn("Cockpit Part", disabled=True),
                "Wiring_Part": st.column_config.TextColumn("Wiring Part", disabled=True),
                "Engine_Stock_After": st.column_config.NumberColumn("Eng Stock After", disabled=True),
                "Cockpit_Stock_After": st.column_config.NumberColumn("CK Stock After", disabled=True),
                "Wiring_Stock_After": st.column_config.NumberColumn("WH Stock After", disabled=True),
            },
        )
        
        # Detect and persist planner edits
        original_display_tcf2 = filtered_df_tcf2[display_cols].reset_index(drop=True)
        edited_display_tcf2 = edited_tcf2.reset_index(drop=True)
        if not original_display_tcf2.equals(edited_display_tcf2):
            for i in range(len(edited_display_tcf2)):
                orig_status = str(original_display_tcf2.at[i, 'STATUS']) if (i < len(original_display_tcf2) and 'STATUS' in original_display_tcf2.columns) else ''
                new_status = str(edited_display_tcf2.at[i, 'STATUS']) if 'STATUS' in edited_display_tcf2.columns else ''
                orig_reason = str(original_display_tcf2.at[i, 'BLOCKING_REASON']) if (i < len(original_display_tcf2) and 'BLOCKING_REASON' in original_display_tcf2.columns) else ''
                new_reason = str(edited_display_tcf2.at[i, 'BLOCKING_REASON']) if 'BLOCKING_REASON' in edited_display_tcf2.columns else ''
                if orig_status != new_status or orig_reason != new_reason:
                    biw_key = str(edited_display_tcf2.at[i, 'BIW NUMBER'])
                    st.session_state.tcf2_manual_overrides[biw_key] = {
                        'status': new_status,
                        'reason': new_reason
                    }
            st.toast("✅ Planner override saved!", icon="✏️")
            filtered_df_tcf2.update(edited_tcf2)
        
        dl_cols_tcf2 = st.columns(2)
        
        # Ready to TCF Excel
        ready_df_tcf2 = filtered_df_tcf2[filtered_df_tcf2['STATUS'] == '✅ Ready for TCF'][display_cols].copy() if 'STATUS' in filtered_df_tcf2.columns else pd.DataFrame()
        if not ready_df_tcf2.empty:
            buf_ready2 = io.BytesIO()
            with pd.ExcelWriter(buf_ready2, engine='openpyxl') as writer:
                ready_df_tcf2.to_excel(writer, index=False, sheet_name='Ready to TCF2')
                ws = writer.sheets['Ready to TCF2']
                hdr_fill = PatternFill(start_color='D8F3E5', end_color='D8F3E5', fill_type='solid')
                hdr_font = Font(name='Calibri', size=11, bold=True, color='1B4D32')
                thin_b = Border(left=Side(style='thin', color='BFBFBF'), right=Side(style='thin', color='BFBFBF'), top=Side(style='thin', color='BFBFBF'), bottom=Side(style='thin', color='BFBFBF'))
                for c in range(1, len(ready_df_tcf2.columns) + 1):
                    cell = ws.cell(row=1, column=c)
                    cell.font = hdr_font
                    cell.fill = hdr_fill
                    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
                    cell.border = thin_b
                for r in range(2, len(ready_df_tcf2) + 2):
                    for c in range(1, len(ready_df_tcf2.columns) + 1):
                        cell = ws.cell(row=r, column=c)
                        cell.border = thin_b
                        cell.alignment = Alignment(horizontal='center', vertical='center')
                        cell.font = Font(name='Calibri', size=10)
                for col in ws.columns:
                    max_len = max(len(str(cell.value or '')) for cell in col)
                    ws.column_dimensions[openpyxl.utils.get_column_letter(col[0].column)].width = max(max_len + 3, 12)
                    
                ready_full_2 = filtered_df_tcf2[filtered_df_tcf2['STATUS'] == '✅ Ready for TCF'].copy()
                vc_s2 = ready_full_2['VEHICLE CODE'] if 'VEHICLE CODE' in ready_full_2.columns else (ready_full_2['VC'] if 'VC' in ready_full_2.columns else pd.Series('', index=ready_full_2.index))
                ready_full_2['Short VC'] = vc_s2.astype(str).str.strip().str[:9]
                pivot_df_2 = ready_full_2.groupby('Short VC').size().reset_index(name='Count')
                tot_row_2 = pd.DataFrame([{'Short VC': 'Total', 'Count': pivot_df_2['Count'].sum()}])
                pivot_full_2 = pd.concat([pivot_df_2, tot_row_2], ignore_index=True)
                
                pivot_full_2.to_excel(writer, index=False, sheet_name='ready to upload')
                ws2_2 = writer.sheets['ready to upload']
                for c in range(1, len(pivot_full_2.columns) + 1):
                    cell = ws2_2.cell(row=1, column=c)
                    cell.font = hdr_font
                    cell.fill = hdr_fill
                    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
                    cell.border = thin_b
                for r in range(2, len(pivot_full_2) + 2):
                    is_tot = (r == len(pivot_full_2) + 1)
                    for c in range(1, len(pivot_full_2.columns) + 1):
                        cell = ws2_2.cell(row=r, column=c)
                        cell.border = thin_b
                        cell.alignment = Alignment(horizontal='center', vertical='center')
                        if is_tot:
                            cell.font = Font(name='Calibri', size=11, bold=True)
                            cell.fill = hdr_fill
                        else:
                            cell.font = Font(name='Calibri', size=10)
                for col in ws2_2.columns:
                    max_len = max(len(str(cell.value or '')) for cell in col)
                    ws2_2.column_dimensions[openpyxl.utils.get_column_letter(col[0].column)].width = max(max_len + 3, 15)
            with dl_cols_tcf2[0]:
                st.download_button(
                    label=f"📥 Ready to TCF2 ({len(ready_df_tcf2)} cabs)",
                    data=buf_ready2.getvalue(),
                    file_name="TCF2_Ready_to_Build.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    key="dl_ready_tcf2"
                )
        
        blocked_df_tcf2 = filtered_df_tcf2[filtered_df_tcf2['STATUS'].isin(['🚫 Blocked', '⚠️ PBS Hold'])][display_cols].copy() if 'STATUS' in filtered_df_tcf2.columns else pd.DataFrame()
        if not blocked_df_tcf2.empty:
            buf_blocked2 = io.BytesIO()
            with pd.ExcelWriter(buf_blocked2, engine='openpyxl') as writer:
                blocked_df_tcf2.to_excel(writer, index=False, sheet_name='Blocked TCF2')
                ws = writer.sheets['Blocked TCF2']
                hdr_fill = PatternFill(start_color='FFD1D1', end_color='FFD1D1', fill_type='solid')
                hdr_font = Font(name='Calibri', size=11, bold=True, color='5C1D1B')
                thin_b = Border(left=Side(style='thin', color='BFBFBF'), right=Side(style='thin', color='BFBFBF'), top=Side(style='thin', color='BFBFBF'), bottom=Side(style='thin', color='BFBFBF'))
                for c in range(1, len(blocked_df_tcf2.columns) + 1):
                    cell = ws.cell(row=1, column=c)
                    cell.font = hdr_font
                    cell.fill = hdr_fill
                    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
                    cell.border = thin_b
                for r in range(2, len(blocked_df_tcf2) + 2):
                    for c in range(1, len(blocked_df_tcf2.columns) + 1):
                        cell = ws.cell(row=r, column=c)
                        cell.border = thin_b
                        cell.alignment = Alignment(horizontal='center', vertical='center')
                        cell.font = Font(name='Calibri', size=10)
                        if len(display_cols) >= 6 and c == display_cols.index('BLOCKING_REASON') + 1:
                            cell.fill = PatternFill(start_color='FFF0F0', end_color='FFF0F0', fill_type='solid')
                for col in ws.columns:
                    max_len = max(len(str(cell.value or '')) for cell in col)
                    ws.column_dimensions[openpyxl.utils.get_column_letter(col[0].column)].width = max(max_len + 3, 12)
            with dl_cols_tcf2[1]:
                st.download_button(
                    label=f"📥 Blocked with Reason ({len(blocked_df_tcf2)} cabs)",
                    data=buf_blocked2.getvalue(),
                    file_name="TCF2_Blocked_with_Reason.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    key="dl_blocked_tcf2"
                )
